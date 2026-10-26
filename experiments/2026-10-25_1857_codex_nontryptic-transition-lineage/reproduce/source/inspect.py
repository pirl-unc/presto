"""Inspect pinned original transition tables with two independent cell readers."""

import argparse
import hashlib
import json
import posixpath
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--requests", type=Path, required=True)
    request = json.loads(parser.parse_args().requests.read_text())
    root = Path(__file__).resolve().parents[1]
    config = json.loads((root / "reproduce/launch.json").read_text())
    data = Path(config["data_directory"])
    inventory = json.loads((data / "files-page0.json").read_text())
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    outputs = []
    for name in request["files"]:
        source = data / name
        content = source.read_bytes()
        candidates = [f for f in inventory if f["fileName"] == name]
        assert len(candidates) == 1
        entry = candidates[0]
        assert entry["fileSizeBytes"] == len(content)
        receipt = json.loads((data / (name + ".receipt.json")).read_text())
        assert receipt["status"] == "retrieved_validated_payload"
        assert hashlib.sha256(content).hexdigest() == receipt["sha256"]
        book = openpyxl.load_workbook(source, data_only=True)
        formula_book = openpyxl.load_workbook(source, data_only=False)
        with ZipFile(source) as z:
            strings = (
                [
                    "".join(n.itertext())
                    for n in ET.fromstring(z.read("xl/sharedStrings.xml")).findall(
                        "s:si", ns
                    )
                ]
                if "xl/sharedStrings.xml" in z.namelist()
                else []
            )
            relations = {
                n.attrib["Id"]: n.attrib["Target"]
                for n in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            }
            for sheet in ET.fromstring(z.read("xl/workbook.xml")).findall(
                "s:sheets/s:sheet", ns
            ):
                title = sheet.attrib["name"]
                target = relations[
                    sheet.attrib[
                        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                    ]
                ]
                member = (
                    target.lstrip("/")
                    if target.startswith("/")
                    else posixpath.normpath(posixpath.join("xl", target))
                )
                records = []
                kinds = Counter()
                formulas = []
                for cell in ET.fromstring(z.read(member)).findall(
                    ".//s:sheetData/s:row/s:c", ns
                ):
                    coord = cell.attrib["r"]
                    kind = cell.attrib.get("t", "n")
                    raw = cell.findtext("s:v", namespaces=ns)
                    if kind == "s":
                        value = strings[int(raw)] if raw is not None else None
                    elif kind == "inlineStr":
                        value = "".join(
                            n.text or "" for n in cell.findall(".//s:t", ns)
                        )
                    elif raw is None:
                        value = None
                    elif kind == "b":
                        value = raw == "1"
                    elif kind in ("str", "e"):
                        value = raw
                    else:
                        value = float(raw)
                    assert book[title][coord].value == value, (
                        name,
                        title,
                        coord,
                        value,
                        book[title][coord].value,
                    )
                    formula = cell.find("s:f", ns)
                    if formula is not None:
                        assert formula_book[title][coord].data_type == "f"
                        formulas.append(
                            {"cell": coord, "formula": formula.text, "cached": raw}
                        )
                    records.append(
                        {
                            "cell": coord,
                            "type": kind,
                            "style": cell.attrib.get("s"),
                            "raw": raw,
                            "value": value,
                        }
                    )
                    kinds[kind] += 1
                preview = []
                for row in book[title].iter_rows():
                    values = [
                        {"cell": c.coordinate, "value": c.value}
                        for c in row
                        if c.value is not None
                    ]
                    if values:
                        preview.append(values)
                    if len(preview) == 8:
                        break
                outputs.append(
                    {
                        "source_file": name,
                        "sheet": title,
                        "dimensions": book[title].calculate_dimension(),
                        "source_cell_count": len(records),
                        "cell_types": dict(kinds),
                        "literal_cells_sha256": hashlib.sha256(
                            json.dumps(records, sort_keys=True).encode()
                        ).hexdigest(),
                        "source_sha256": receipt["sha256"],
                        "source_bytes": len(content),
                        "opaque_inventory_checksum": entry["checksum"],
                        "matches_computed_sha1": hashlib.sha1(content).hexdigest()
                        == entry["checksum"],
                        "checksum_algorithm_declared_by_inventory": None,
                        "merged_cells": str(book[title].merged_cells),
                        "formulas": formulas,
                        "first_nonempty_rows": preview,
                    }
                )
        book.close()
        formula_book.close()
    result = {
        "status": "literal_cells_verified_semantics_pending",
        "n_workbooks": len(request["files"]),
        "n_sheets": len(outputs),
        "n_source_cells_verified": sum(x["source_cell_count"] for x in outputs),
        "worksheets": outputs,
        "model_fits": 0,
    }
    with (root / "results/transition-inventory.json").open("x") as out:
        json.dump(result, out, indent=2, allow_nan=False)
        out.write("\n")
    print(json.dumps({k: v for k, v in result.items() if k != "worksheets"}))


if __name__ == "__main__":
    main()
