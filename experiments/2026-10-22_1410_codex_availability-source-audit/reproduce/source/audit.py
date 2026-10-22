"""Read-only source inventory; not a Hitlist catalog or training-data adapter."""

import argparse
import collections
import csv
import hashlib
import json
import math
from pathlib import Path
import platform
import statistics
import time
import zipfile

import datacache
import openpyxl
from openpyxl.utils import get_column_letter


FAMILY = Path(__file__).resolve().parents[2 if Path(__file__).parent.name == "source" else 1]
REPO = FAMILY.parents[1]


def write_json(path, obj):
    path.write_text(json.dumps(obj, indent=2, sort_keys=True, allow_nan=False) + "\n")


def workbook(path, table):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, path
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    assert book.sheetnames == [f"Supplementary Table {table}"]
    sheet = book.active
    rows = list(sheet.values)
    assert rows[4][0] == "Sequence"
    protein_column = rows[4].index("Protein")
    # Published quantitative block is the leftmost block, with channel labels
    # on row 4 and sample labels on row 5. Later panels are distinct results.
    quantity_columns = []
    for column in range(1, protein_column):
        if rows[3][column] is None:
            break
        quantity_columns.append(column)
    primary = [(index + 6, row) for index, row in enumerate(rows[5:]) if row[0] is not None]
    assert len(primary) == len(rows) - 5, "Unexpected blanks require review"
    sequences = [row[0] for _, row in primary]
    parents = [row[protein_column] for _, row in primary]
    channels = []
    for column in quantity_columns:
        values = [row[column] for _, row in primary]
        numeric = [x for x in values if isinstance(x, (int, float)) and not isinstance(x, bool)]
        assert all(math.isfinite(x) for x in numeric)
        channels.append({
            "column": get_column_letter(column + 1),
            "channel_cell": f"{get_column_letter(column + 1)}4",
            "sample_cell": f"{get_column_letter(column + 1)}5",
            "source_channel_label": rows[3][column],
            "source_sample_label": rows[4][column],
            "numeric_cells": len(numeric),
            "blank_cells": sum(x is None for x in values),
            "nonnumeric_nonblank_cells": sum(x is not None for x in values) - len(numeric),
            "zero_cells": sum(x == 0 for x in numeric),
            "negative_cells": sum(x < 0 for x in numeric),
            "sum": math.fsum(numeric),
            "minimum": min(numeric),
            "maximum": max(numeric),
        })
    blocks = []
    for column, value in enumerate(rows[4]):
        if value == "Sequence":
            block_rows = [(index + 6, row[column]) for index, row in enumerate(rows[5:])
                          if row[column] is not None]
            blocks.append({
                "sequence_header_cell": f"{get_column_letter(column + 1)}5",
                "populated_sequence_rows": len(block_rows),
                "distinct_sequence_labels": len({value for _, value in block_rows}),
                "same_row_sequence_disagreements_with_primary": sum(
                    value != rows[index - 1][0] for index, value in block_rows),
            })
    sums = [c["sum"] for c in channels]
    result = {
        "table": table, "sheet": sheet.title,
        "reported_sheet_shape": [sheet.max_row, sheet.max_column],
        "primary_block_range": f"A5:{get_column_letter(protein_column + 1)}{len(rows)}",
        "primary_rows": len(primary),
        "distinct_sequence_labels": len(set(sequences)),
        "distinct_nonblank_parent_labels": len(set(parents) - {None, ""}),
        "blank_parent_labels": sum(x in (None, "") for x in parents),
        "sequence_length_min_max": [min(map(len, sequences)), max(map(len, sequences))],
        "quantity_cells": len(primary) * len(channels),
        "channels": channels,
        "channel_total_cv": statistics.pstdev(sums) / statistics.mean(sums),
        "independently_keyed_sequence_blocks": blocks,
        "formula_cells": sum(isinstance(x, str) and x.startswith("=") for row in rows for x in row),
        "header_cells": [
            {"cell": f"{get_column_letter(c + 1)}{r + 1}", "value": value}
            for r, row in enumerate(rows[:5]) for c, value in enumerate(row) if value is not None
        ],
    }
    book.close()
    return result, set(parents) - {None, ""}


def sdrf(path, deposit_files):
    with path.open(newline="") as handle:
        all_rows = list(csv.reader(handle, delimiter="\t"))
    header, rows = all_rows[0], all_rows[1:]
    assert all(len(row) == len(header) for row in rows)
    # Columns with repeated names are legal source content, not dictionary keys.
    columns = [{"column_index_1based": i + 1, "name": name,
                "values": dict(collections.Counter(row[i] for row in rows))}
               for i, name in enumerate(header)]
    file_column = header.index("comment[data file]")
    observed = {row[file_column] for row in rows}
    raw = {x["fileName"] for x in deposit_files if x["fileCategory"]["value"] == "RAW"}
    assert observed <= raw
    return {
        "rows": len(rows), "columns": len(header),
        "duplicate_column_names": {k: v for k, v in collections.Counter(header).items() if v > 1},
        "column_values": columns,
        "distinct_source_names": len({row[header.index("source name")] for row in rows}),
        "annotated_raw_files": sorted(observed),
        "unannotated_raw_files": sorted(raw - observed),
        "deposit_file_categories": dict(collections.Counter(
            x["fileCategory"]["value"] for x in deposit_files)),
        "full_sdrf_validator_run": False,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-paths", type=Path, help="Optional local originals; hashes still required")
    parser.add_argument("--output", type=Path, default=FAMILY / "results")
    args = parser.parse_args()
    start = time.perf_counter()
    inventory_path = FAMILY / "reproduce" / "input_assets.json"
    inventory = json.loads(inventory_path.read_text())
    assets = inventory["assets"]
    bundle_id = hashlib.sha256(json.dumps(assets, sort_keys=True).encode()).hexdigest()
    snapshot_root = REPO / "data/availability_sources/snapshots/2026-10-22"
    local_paths = {name: str(snapshot_root / name) for name in assets
                   if (snapshot_root / name).is_file()}
    if args.source_paths:
        local_paths.update(json.loads(args.source_paths.read_text()))
    paths = datacache.install_bundle(
        REPO / "data" / "availability_sources" / bundle_id,
        assets, source_paths=local_paths, download_options={"show_progress": False},
    )
    args.output.mkdir(parents=True, exist_ok=True)
    tables, parent_sets = [], {}
    for number in range(7, 13):
        result, parents = workbook(Path(paths[f"PXD057588/table{number}.xlsx"]), number)
        tables.append(result)
        parent_sets[number] = parents
    deposit = json.loads(Path(paths["PXD057588/files.json"]).read_text())
    sdrf_result = sdrf(Path(paths["PXD057588/community.sdrf.tsv"]), deposit)
    table_by_number = {x["table"]: x for x in tables}
    map9 = {str(c["source_channel_label"]): c["source_sample_label"]
            for c in table_by_number[9]["channels"]}
    map11 = {str(c["source_channel_label"]): c["source_sample_label"]
             for c in table_by_number[11]["channels"]}
    sdrf_labels = set(next(c["values"] for c in sdrf_result["column_values"]
                          if c["name"] == "comment[label]"))
    comparison = {
        "table9_table11_channel_label_to_sample_label": [
            {"channel": key, "table9_sample": map9[key], "table11_sample": map11.get(key),
             "same_sample_label": map9[key] == map11.get(key)} for key in map9
        ],
        "table9_table11_shared_parent_labels": len(parent_sets[9] & parent_sets[11]),
        "table9_parent_labels_without_table11_exact_label_match": len(parent_sets[9] - parent_sets[11]),
        "table12_sdrf_channel_label_compatibility": [
            {"source_channel_label": c["source_channel_label"],
             "table12_sample_label": c["source_sample_label"],
             "prefixed_label_present_in_sdrf": "TMT" + str(c["source_channel_label"]) in sdrf_labels}
            for c in table_by_number[12]["channels"]
        ],
        "interpretation": "Label overlaps only; no verified specimen/aliquot or sequence-to-reference join.",
        "training_admission": "deferred_pending_Hitlist_curation_and_endpoint_contract",
    }
    write_json(args.output / "workbook_inventory.json", tables)
    write_json(args.output / "sdrf_inventory.json", sdrf_result)
    write_json(args.output / "cross_assay_checks.json", comparison)
    with (args.output / "table_summary.csv").open("w", newline="") as handle:
        fields = ["table", "primary_rows", "distinct_sequence_labels",
                  "distinct_nonblank_parent_labels", "quantity_cells", "channel_total_cv"]
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows({key: result[key] for key in fields} for result in tables)
    write_json(args.output / "runtime.json", {
        "seconds": time.perf_counter() - start, "python": platform.python_version(),
        "platform": platform.platform(), "openpyxl": openpyxl.__version__,
        "datacache": datacache.__version__, "gpu": None, "training": False,
        "bundle_id": bundle_id, "asset_count": len(assets),
        "total_input_bytes": sum(x["size"] for x in assets.values()),
        "input_assets_sha256": hashlib.sha256(inventory_path.read_bytes()).hexdigest(),
    })
    print(json.dumps({"tables": len(tables), "sdrf_rows": sdrf_result["rows"],
                      "raw_files_annotated": len(sdrf_result["annotated_raw_files"]),
                      "bundle_id": bundle_id}))


if __name__ == "__main__":
    main()
