"""Replay recipe admission from original TSV/XLSX and independent reference hits."""

import csv
import gzip
import hashlib
import json
import math
import resource
import statistics
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

from Bio import SeqIO
from openpyxl import load_workbook

FAMILY = Path(__file__).resolve().parents[1]
SOURCE = Path("/private/tmp/hitlist-v2-mapp-reference-second")
ASSETS = SOURCE / "external/pxd009092/assets"
PRIOR = FAMILY.parent / "2026-10-22_1845_codex_mapp-parent-reference/results"
RID = "reviewed_2026_03_isoforms"
AA = set("ACDEFGHIKLMNPQRSTVWY")
CONDITIONS = ("UT", "TNF+IFN")


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def jsonl(path):
    with Path(path).open() as stream:
        return [json.loads(line) for line in stream]


def tsv(name):
    with (ASSETS / "txt" / f"{name}.txt").open(newline="") as stream:
        yield from csv.DictReader(stream, delimiter="\t")


def main():
    tick = time.monotonic()
    output = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    assert not destination.exists()
    manifest = json.loads((output / "manifest.json").read_text())
    assert (
        sha(SOURCE / "manifest.json")
        == manifest["source_manifest_sha256"]
        == "55f7a3d50df5d162e55bf1d66fbcf413a35fc927fd08c266f90d981068a053c2"
    )
    source_manifest = json.loads((SOURCE / "manifest.json").read_text())
    for name, record in source_manifest["files"].items():
        assert (
            sha(SOURCE / name) == record["sha256"]
            and (SOURCE / name).stat().st_size == record["size_bytes"]
        )
    assert {p.name for p in output.iterdir()} == set(manifest["files"]) | {
        "manifest.json"
    }
    for name, record in manifest["files"].items():
        assert (
            sha(output / name) == record["sha256"]
            and (output / name).stat().st_size == record["size_bytes"]
        )
    recipe = json.loads((output / "recipe.json").read_text())
    minimum = recipe["selection"]["min_eligible_peptides_per_parent"]
    assert minimum in (2, 3)
    peptides = {p["id"]: p for p in tsv("peptides")}
    assert len(peptides) == 10673
    runs = defaultdict(list)
    for r in tsv("summary"):
        if r["Experiment"]:
            condition = r["Experiment"].rsplit("_", 2)[1]
            runs[condition].append(r["Experiment"])
    assert all(len(runs[c]) == 3 for c in CONDITIONS)
    direct, evidence = defaultdict(set), defaultdict(list)
    for r in tsv("evidence"):
        condition = r["Experiment"].rsplit("_", 2)[1]
        if condition in CONDITIONS:
            evidence[r["Peptide ID"], condition].append("pxd009092:evidence:" + r["id"])
            if int(r["MS/MS Count"] or 0) > 0:
                direct[r["Peptide ID"], condition].add(r["Experiment"])
    publisher = {}
    wb = load_workbook(
        ASSETS / "publisher/suppdata2.xlsx", read_only=True, data_only=True
    )
    try:
        for index, row in enumerate(
            wb["Sheet1"].iter_rows(min_row=5, values_only=True), 5
        ):
            if not row[0]:
                continue
            for col, condition in ((2, "UT"), (3, "TNF+IFN")):
                if row[col] is not None:
                    publisher[row[0], condition] = (
                        float(row[col]),
                        f"publisher/suppdata2.xlsx:Sheet1:{chr(65 + col)}{index}",
                    )
    finally:
        wb.close()
    sequences = {
        r.id.split("|")[1]: str(r.seq)
        for r in SeqIO.parse(
            ASSETS / "references/human-reviewed-2026_03-isoforms.fasta", "fasta"
        )
    }
    sequence_hashes = {
        a: hashlib.sha256(s.encode()).hexdigest() for a, s in sequences.items()
    }
    by_hash = {h: sequences[a] for a, h in sequence_hashes.items()}
    reference_aliases = defaultdict(list)
    for acc, h in sequence_hashes.items():
        reference_aliases[h].append(RID + ":" + acc)
    old_hashes = json.loads((PRIOR / "output_hashes.json").read_text())
    matches_path = PRIOR / f"{RID}_matches.csv.gz"
    assert sha(matches_path) == old_hashes[matches_path.name]
    matches = defaultdict(list)
    with gzip.open(matches_path, "rt", newline="") as stream:
        for r in csv.DictReader(stream):
            pid, acc, start = r["peptide_id"], r["accession"], int(r["start1"]) - 1
            assert sequences[acc][start : int(r["end1"])] == peptides[pid]["Sequence"]
            assert r["sequence_sha256"] == sequence_hashes[acc]
            matches[pid].append((acc, start))
    # Independent bipartite traversal on full-sequence identity and original
    # source protein groups, including every excluded/unmatched peptide.
    neighbors, owners = defaultdict(set), defaultdict(set)
    for pid, p in peptides.items():
        links = {("sequence", sequence_hashes[a]) for a, _ in matches[pid]}
        links |= {("group", g) for g in p["Protein group IDs"].split(";") if g}
        neighbors[pid] = links
        for link in links:
            owners[link].add(pid)
    components = {}
    for pid in peptides:
        if pid in components:
            continue
        pending = [pid]
        members = {pid}
        visited = set()
        while pending:
            p = pending.pop()
            for link in neighbors[p] - visited:
                visited.add(link)
                for other in owners[link] - members:
                    members.add(other)
                    pending.append(other)
        label = min("pxd009092:peptide:" + p for p in members)
        for p in members:
            components[p] = label
    expected = {}
    groups = defaultdict(list)
    for pid, p in peptides.items():
        failures = []
        checks = {
            "source_reverse": p["Reverse"] == "+",
            "source_contaminant": p["Potential contaminant"] == "+",
            "noncanonical_peptide": not p["Sequence"] or not set(p["Sequence"]) <= AA,
            "ambiguous_source_group": len(
                set(filter(None, p["Protein group IDs"].split(";")))
            )
            != 1,
        }
        for c in CONDITIONS:
            values = [
                float(p["Intensity " + r.replace("T+I_", "T+I-MAPPbeadFilt_")])
                for r in runs[c]
            ]
            checks["insufficient_positive_runs:" + c] = sum(v > 0 for v in values) < 2
            checks["missing_direct_msms:" + c] = not direct[pid, c]
            if (p["Sequence"], c) in publisher:
                assert math.isclose(
                    10 ** publisher[p["Sequence"], c][0],
                    statistics.median(values),
                    rel_tol=1e-4,
                    abs_tol=1,
                )
        structural = not any(checks.values())
        for label, failed in checks.items():
            if failed:
                failures.append(label)
        pep = float(p["PEP"]) if p["PEP"] else float("nan")
        if not math.isfinite(pep):
            failures.append("peptide_confidence_unreported")
        elif not 0 <= pep <= 0.01:
            failures.append("peptide_pep_above_threshold")
        for c in CONDITIONS:
            if (p["Sequence"], c) not in publisher:
                failures.append("missing_nonimputed_quantity:" + c)
        full = {sequence_hashes[a] for a, _ in matches[pid]}
        if not full:
            failures.append("reference_unmapped")
        elif len(full) > 1:
            failures.append("multiple_parent_sequences")
        seq = sequences.get(p["Leading razor protein"])
        leading = False
        if (
            seq is not None
            and p["Start position"].isdigit()
            and p["End position"].isdigit()
        ):
            start, end = int(p["Start position"]), int(p["End position"])
            leading = (
                1 <= start <= end <= len(seq)
                and seq[start - 1 : end] == p["Sequence"]
                and p["Amino acid before"] == (seq[start - 2] if start > 1 else "-")
                and p["Amino acid after"] == (seq[end] if end < len(seq) else "-")
            )
        if not leading:
            failures.append("source_leading_coordinate_or_flank_unconfirmed")
        parent = next(iter(full)) if len(full) == 1 else None
        if parent is not None and not set(by_hash[parent]) <= AA:
            failures.append("noncanonical_parent")
        bucket = (
            int.from_bytes(
                hashlib.sha256(
                    ("mapp_captured_composition_v1:" + components[pid]).encode()
                ).digest()[:8],
                "big",
            )
            % 100
        )
        split = "train" if bucket < 70 else ("validation" if bucket < 85 else "test")
        expected[pid] = {
            "reasons": failures,
            "structural_support": structural,
            "parent_sequence_sha256": parent,
            "n_reference_sequences": len(full),
            "n_reference_occurrences": len(matches[pid]),
            "component_id": components[pid],
            "split": split,
        }
        if not failures:
            groups[parent].append(pid)
    for parent, pids in groups.items():
        if len(pids) < minimum:
            for pid in pids:
                expected[pid]["reasons"].append(
                    f"fewer_than_{'three' if minimum == 3 else 'two'}_eligible_peptides_per_parent"
                )
    actual_ledger = jsonl(output / "peptide_dispositions.jsonl")
    assert len(actual_ledger) == len(expected)
    eligible = {pid for pid, r in expected.items() if not r["reasons"]}
    for r in actual_ledger:
        pid = r["peptide_row_id"].rsplit(":", 1)[1]
        wanted = expected[pid]
        assert set(r["reasons"]) == set(wanted["reasons"]), (pid, r, wanted)
        for k, v in wanted.items():
            if k != "reasons":
                assert r[k] == v, (pid, k)
        assert r["status"] == ("admitted" if pid in eligible else "excluded")
        assert r["source_record_id"] == f"pxd009092:txt/peptides.txt:row:{int(pid) + 2}"
    parents = jsonl(output / "parents.jsonl")
    assert len(parents) == sum(len(pids) >= minimum for pids in groups.values())
    for r in parents:
        h = r["parent_sequence_sha256"]
        assert r["sequence"] == by_hash[h]
        assert r["reference_record_ids"] == sorted(reference_aliases[h])
        assert r["component_id"] == components[groups[h][0]]
        assert r["split"] == expected[groups[h][0]]["split"]
    examples = jsonl(output / "examples.jsonl")
    assert len(examples) == 2 * len(parents)
    covered = Counter()
    distributions = Counter()
    for r in examples:
        h, c = r["parent_sequence_sha256"], r["condition"]
        pids = groups[h]
        assert r["context"] == [float(c == "TNF+IFN")]
        assert (
            r["split"] == expected[pids[0]]["split"]
            and r["component_id"] == components[pids[0]]
        )
        assert (
            r["measured_species"] == "Homo sapiens" and r["measured_cell"] == "HEK293"
        )
        assert r["material_link_precision"] == "acquisition_run_unresolved_specimen"
        assert {q["peptide_row_id"].rsplit(":", 1)[1] for q in r["candidates"]} == set(
            pids
        )
        values = [publisher[peptides[pid]["Sequence"], c][0] for pid in pids]
        total = sum(10 ** (v - max(values)) for v in values)
        for q in r["candidates"]:
            pid = q["peptide_row_id"].rsplit(":", 1)[1]
            value, measurement = publisher[peptides[pid]["Sequence"], c]
            assert (
                q["source_log10_intensity"] == value
                and q["measurement_id"] == measurement
            )
            assert q["sequence"] == peptides[pid]["Sequence"]
            assert q["start_positions_0based"] == sorted(
                {start for _, start in matches[pid]}
            )
            target = 10 ** (value - max(values)) / total
            assert math.isclose(q["target_probability"], target, rel_tol=1e-12)
            assert math.isclose(
                q["target_log_probability"],
                math.log(target),
                rel_tol=1e-12,
                abs_tol=1e-12,
            )
            assert math.isclose(
                q["target_centered_log10"],
                value - statistics.mean(values),
                abs_tol=1e-12,
            )
            covered[pid, c] += 1
        distributions[r["split"]] += 1
    assert set(covered) == {(pid, c) for pid in eligible for c in CONDITIONS} and set(
        covered.values()
    ) == {1}
    contributors = jsonl(output / "contributors.jsonl")
    assert len(contributors) == len(covered)
    for r in contributors:
        pid, c = r["peptide_row_id"].rsplit(":", 1)[1], r["condition"]
        assert (
            r["measurement_id"]
            == publisher[peptides[pid]["Sequence"], c][1]
            == r["source_cell"]
        )
        assert r["run_measurement_ids"] == sorted(
            f"pxd009092:peptide:{pid}:run:{run}:bead_filtered" for run in runs[c]
        )
        assert r["channel_ids"] == sorted("pxd009092:run:" + run for run in runs[c])
        assert r["direct_msms_channel_ids"] == sorted(
            "pxd009092:run:" + run for run in direct[pid, c]
        )
        assert r["evidence_ids"] == sorted(evidence[pid, c])
    summary = json.loads((output / "summary.json").read_text())
    assert (
        summary["n_structurally_supported_peptides"]
        == sum(r["structural_support"] for r in expected.values())
        == 2272
    )
    assert summary["n_admitted_peptides"] == len(eligible)
    assert summary["n_parents_by_split"] == dict(Counter(p["split"] for p in parents))
    assert summary["exclusion_reasons"] == dict(
        Counter(reason for r in expected.values() for reason in r["reasons"])
    )
    assert summary["n_all_source_components"] == len(set(components.values()))
    assert summary["largest_source_component_n_peptides"] == max(
        Counter(components.values()).values()
    )
    report = {
        "status": "passed",
        "recipe_id": recipe["recipe_id"],
        "minimum_peptides": minimum,
        "source_manifest_sha256": sha(SOURCE / "manifest.json"),
        "recipe_manifest_sha256": sha(output / "manifest.json"),
        "n_source_peptides": len(expected),
        "n_admitted_peptides": len(eligible),
        "n_parents": len(parents),
        "n_distributions_by_split": dict(distributions),
        "n_contributors": len(contributors),
        "n_all_source_components": len(set(components.values())),
        "seconds": time.monotonic() - tick,
        "peak_process_rss_bytes": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss,
        "scope": "All original TSV/XLSX admission, quantities, contributors, independent reference hits, all-source graph and fixed holdouts replayed without importing the producer.",
    }
    destination.write_text(json.dumps(report, indent=2, sort_keys=True) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
