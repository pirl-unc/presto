"""Reconcile original publisher cells and deposited MAPP background filtering."""

from collections import Counter
import csv
from datetime import datetime, timezone
import gzip
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import time

import datacache
import numpy as np
import openpyxl
import pandas as pd

FAMILY = Path(__file__).resolve().parents[1]
ROOT = FAMILY.parents[1]
PREVIOUS = ROOT / "experiments/2026-10-22_1630_codex_mapp-source-admission"
RESULTS = FAMILY / "results"
CONDITIONS = ("UT", "TNF+IFN", "CHX", "NO-DSP", "NO-INH")


def digest(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(8 * 1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def dump(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n")


def csv_dump(path, rows):
    with gzip.open(path, "wt", newline="") as stream:
        writer = None
        for row in rows:
            if writer is None:
                writer = csv.DictWriter(stream, fieldnames=list(row), lineterminator="\n")
                writer.writeheader()
            writer.writerow(row)


def source_paths():
    receipt = json.loads((FAMILY / "reproduce/publisher_assets.json").read_text())
    key = hashlib.sha256(json.dumps(receipt["assets"], sort_keys=True).encode()).hexdigest()
    paths = datacache.install_bundle(ROOT / "data/availability_sources" / key,
        receipt["assets"], source_paths=receipt["seed_paths"])
    previous = json.loads((PREVIOUS / "results/input_paths.json").read_text())
    hashes = json.loads((PREVIOUS / "reproduce/archive_members.json").read_text())
    members = {name: ROOT / path for name, path in previous["archive_members"].items()}
    for name, path in members.items():
        assert digest(path) == hashes[name]["sha256"]
    return paths, members


def main():
    tick = time.monotonic()
    RESULTS.mkdir(parents=True, exist_ok=True)
    assert not (RESULTS / "summary.json").exists(), "preserve completed audit"
    dump(FAMILY / "reproduce/launch.json", {
        "agent_model": "Codex / GPT-6", "timestamp": datetime.now(timezone.utc).isoformat(),
        "argv": sys.argv, "executable": sys.executable, "cwd": str(ROOT),
        "git_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], text=True).strip(),
        "git_status": subprocess.check_output(["git", "status", "--porcelain"], text=True),
        "python": sys.version, "numpy": np.__version__, "pandas": pd.__version__,
        "openpyxl": openpyxl.__version__, "model": None, "training": None,
        "synthetic": None, "split": None, "device": "local CPU; source audit only",
        "no_predictive_evaluation_reason": "source quantities and provenance reconciliation only",
        "code_sha256": digest(__file__),
    })
    shutil.copy2(__file__, FAMILY / "reproduce/source/audit.py")
    publisher, members = source_paths()
    dump(RESULTS / "input_paths.json", {"publisher": {k: str(v) for k, v in publisher.items()},
        "archive_members": {k: str(v) for k, v in members.items()}})
    workbooks, inventory = {}, {}
    for name in ("suppdata2.xlsx", "suppdata3.xlsx"):
        workbook = openpyxl.load_workbook(publisher[name], read_only=True, data_only=True)
        workbooks[name] = {}
        inventory[name] = {}
        for sheet in workbook:
            values = list(sheet.values)
            workbooks[name][sheet.title] = values
            inventory[name][sheet.title] = {"rows": sheet.max_row, "columns": sheet.max_column,
                "leading_rows": values[:4], "merged_cells": "not interpreted in read-only mode"}
        workbook.close()
    dump(RESULTS / "workbook_inventory.json", inventory)
    with gzip.open(RESULTS / "original_workbook_cells.json.gz", "wt") as stream:
        json.dump(workbooks, stream, allow_nan=False)
    peptides = pd.read_csv(members["txt/peptides.txt"], sep="\t", dtype=str, keep_default_na=False)
    assert peptides["id"].is_unique and peptides["Sequence"].is_unique
    peptides = peptides.set_index("id", drop=False)
    evidence = pd.read_csv(members["txt/evidence.txt"], sep="\t", dtype=str, keep_default_na=False,
        usecols=["Peptide ID", "Experiment", "Intensity"])
    evidence["value"] = evidence["Intensity"].replace("", "0").astype(float)
    quantities = evidence.groupby(["Peptide ID", "Experiment"])["value"].sum().unstack(fill_value=0)
    quantities = quantities.reindex(peptides.index, fill_value=0)
    columns = {c: c.removeprefix("Intensity ").replace("-MAPPbeadFilt", "")
               for c in peptides.columns if c.startswith("Intensity ")}
    assert len(columns) == 15
    experiments = list(columns.values())
    measured = peptides[list(columns)].astype(float).to_numpy()
    raw = quantities[experiments].to_numpy()
    bead_experiments = sorted(c for c in quantities if "_BEADS_" in c)
    assert len(bead_experiments) == 3
    beads = quantities[bead_experiments].to_numpy()
    condition_indices = {c: [j for j, e in enumerate(experiments) if f"_{c}_" in e]
                         for c in CONDITIONS}
    assert all(len(v) == 3 for v in condition_indices.values())
    rule_names, masks, rule_summaries = [], [], []
    for background_name, aggregate in (("mean", np.mean), ("median", np.median), ("max", np.max)):
        background = aggregate(beads, axis=1)
        for numerator_name in ("run", "condition_mean", "condition_median"):
            if numerator_name == "run":
                numerator = raw
            else:
                numerator = np.empty_like(raw)
                function = np.mean if numerator_name == "condition_mean" else np.median
                for indices in condition_indices.values():
                    numerator[:, indices] = function(raw[:, indices], axis=1)[:, None]
            mask = numerator >= 2 * background[:, None]
            predicted = np.where(mask, raw, 0)
            agrees = np.isclose(predicted, measured, rtol=1e-4, atol=1)
            name = f"{numerator_name}_over_beads_{background_name}"
            rule_names.append(name)
            masks.append(mask)
            rule_summaries.append({"rule": name, "cells": int(raw.size),
                "matching_cells": int(agrees.sum()), "mismatching_cells": int((~agrees).sum()),
                "reported_zero_predicted_positive": int(((measured == 0) & (predicted > 0)).sum()),
                "reported_positive_predicted_zero": int(((measured > 0) & (predicted == 0)).sum()),
                "positive_numerator_exact_boundary": int(((numerator > 0) & (numerator == 2*background[:, None])).sum())})
    np.savez_compressed(RESULTS / "filter_comparisons.npz", peptide_ids=peptides.index.to_numpy(dtype=str),
        experiments=np.asarray(experiments), measured=measured, raw=raw, beads=beads,
        bead_experiments=np.asarray(bead_experiments), rule_names=np.asarray(rule_names),
        masks=np.stack(masks))
    dump(RESULTS / "background_rules.json", rule_summaries)
    source_medians = {}
    for condition, indices in condition_indices.items():
        vals = measured[:, indices]
        source_medians[(condition, "include_zero")] = np.median(vals, axis=1)
        # Preserve an undefined all-zero positive-only median as NaN internally.
        positive = pd.DataFrame(vals).where(vals > 0).median(axis=1).to_numpy()
        source_medians[(condition, "positive_only")] = positive
    original_rows = workbooks["suppdata2.xlsx"]["Sheet1"]
    published = [(i, row) for i, row in enumerate(original_rows, 1) if i >= 5 and row[0] is not None]
    seq_to_pos = {s: i for i, s in enumerate(peptides["Sequence"])}
    seq_counts = Counter(row[0] for _, row in published)
    assert all(count == 1 for count in seq_counts.values()), "retain and diagnose duplicate publisher IDs"
    cell_rows, counters = [], Counter()
    column_summary = []
    for letter, column, encoding in (("C", 2, "log10"), ("D", 3, "log10"),
                                     ("E", 4, "linear"), ("F", 5, "linear")):
        counts = Counter()
        for row_number, row in published:
            sequence, value = row[0], row[column]
            counts["rows"] += 1
            if value is None:
                counts["missing"] += 1
                continue
            assert isinstance(value, (int, float)) and np.isfinite(value)
            counts["numeric"] += 1
            linear = float(10 ** value if encoding == "log10" else value)
            position = seq_to_pos.get(sequence)
            if position is None:
                counts["sequence_not_in_deposit"] += 1
            for (condition, aggregation), source_values in source_medians.items():
                source = float(source_values[position]) if position is not None else float("nan")
                matches = bool(np.isfinite(source) and np.isclose(linear, source, rtol=1e-4, atol=1))
                key = (letter, condition, aggregation)
                counters[(key, "compared")] += int(np.isfinite(source))
                counters[(key, "matched")] += matches
                cell_rows.append({"sheet": "Sheet1", "cell": f"{letter}{row_number}",
                    "sequence": sequence, "gene_label": row[1], "published_value": value,
                    "published_encoding": encoding, "published_linear": linear,
                    "deposited_peptide_id": peptides.iloc[position]["id"] if position is not None else None,
                    "candidate_source_condition": condition, "median_aggregation": aggregation,
                    "source_median": source if np.isfinite(source) else None, "matches": matches})
        column_summary.append({"column": letter, "header": original_rows[2][column],
            "encoding": encoding, "counts": dict(counts)})
    csv_dump(RESULTS / "publisher_cell_comparisons.csv.gz", cell_rows)
    matches = [{"column": k[0], "source_condition": k[1], "aggregation": k[2],
        "compared": counters[(k, "compared")], "matched": counters[(k, "matched")]}
        for k in sorted({key for key, _ in counters})]
    dump(RESULTS / "publisher_column_matches.json", matches)
    seq_ids = [seq_to_pos[s] for s in seq_counts if s in seq_to_pos]
    selected = peptides.iloc[seq_ids]
    data3 = workbooks["suppdata3.xlsx"]["Sheet1"]
    summary = {"publisher_peptide_rows": len(published), "publisher_unique_sequences": len(seq_counts),
        "publisher_sequences_in_deposit": len(seq_ids),
        "publisher_sequences_absent_from_deposit": sorted(set(seq_counts)-set(seq_to_pos)),
        "publisher_reverse_rows": int((selected["Reverse"] == "+").sum()),
        "publisher_contaminant_rows": int((selected["Potential contaminant"] == "+").sum()),
        "columns": column_summary,
        "publisher_protein_summary_rows": len([r for r in data3[2:] if r[0] is not None]),
        "publisher_protein_headers": data3[:2],
        "best_background_rule_by_cell_match": max(rule_summaries, key=lambda r:r["matching_cells"]),
        "background_rule_complete_matches": [r["rule"] for r in rule_summaries if r["mismatching_cells"] == 0],
        "all_training_admission": "pending; no source filters or specimen joins inferred from best approximate matches",
        "seconds": time.monotonic()-tick}
    dump(RESULTS / "summary.json", summary)
    dump(RESULTS / "output_hashes.json", {p.name: digest(p) for p in RESULTS.iterdir()
        if p.is_file() and p.name != "output_hashes.json"})
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
