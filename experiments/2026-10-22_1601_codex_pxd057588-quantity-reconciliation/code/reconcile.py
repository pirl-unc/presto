"""Reconcile published biological quantities; no fitting or dataset admission."""

import collections
import csv
import gzip
import hashlib
import io
import json
from pathlib import Path
import platform
import time

import datacache
import numpy as np
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


FAMILY = Path(__file__).resolve().parents[2 if Path(__file__).parent.name == "source" else 1]
REPO = FAMILY.parents[1]
PREVIOUS = REPO / "experiments/2026-10-22_1410_codex_availability-source-audit"


def save(path, value):
    path.write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n")


def inputs(contract):
    manifest = PREVIOUS / "reproduce/input_assets.json"
    assert hashlib.sha256(manifest.read_bytes()).hexdigest() == contract["input_assets_sha256"]
    assets = json.loads(manifest.read_text())["assets"]
    identity = hashlib.sha256(json.dumps(assets, sort_keys=True).encode()).hexdigest()
    snapshots = REPO / "data/availability_sources/snapshots/2026-10-22"
    paths = datacache.install_bundle(
        REPO / "data/availability_sources" / identity, assets,
        source_paths={name: str(snapshots / name) for name in assets if (snapshots / name).is_file()},
        download_options={"show_progress": False},
    )
    return paths


def load_table(path, spec):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    assert workbook.sheetnames == [f"Supplementary Table {spec['table']}"]
    sheet = workbook.active
    rows = list(sheet.values)
    workbook.close()
    assert len(rows) == spec["n_rows"] + 5 and rows[4][0] == "Sequence"
    assert rows[4][column_index_from_string(spec["parent_column"]) - 1] == "Protein"
    columns = []
    for column in range(1, len(rows[3])):
        if rows[3][column] is None:
            break
        columns.append(column)
    samples = {rows[4][column]: index for index, column in enumerate(columns)}
    assert len(samples) == len(columns)
    values = np.array([[row[c] for c in columns] for row in rows[5:]], dtype=np.float64)
    assert np.isfinite(values).all() and (values > 0).all()
    return rows, columns, samples, values


def errors(values, tolerance):
    absolute = np.abs(values)
    return {
        "n_cells": int(values.size),
        "n_mismatches": int(np.count_nonzero(absolute > tolerance)),
        "maximum_absolute_error_log2": float(absolute.max()),
        "mean_absolute_error_log2": float(absolute.mean()),
        "rmse_log2": float(np.sqrt(np.mean(values ** 2))),
    }


def main():
    start = time.perf_counter()
    contract = json.loads((FAMILY / "reproduce/contract.json").read_text())
    paths = inputs(contract)
    results = FAMILY / "results"
    results.mkdir(exist_ok=True)
    summaries, channels, support, parent_counts = [], [], [], {}
    fields = ["table", "source_row", "sequence", "parent_label", "sample", "intensity_cell",
              "fold_change_cell", "control_columns", "reported_log2fc", "paper_policy_log2fc",
              "wrong_policy_log2fc", "paper_policy_error_log2", "wrong_policy_error_log2"]
    with (results / "cell_comparisons.csv.gz").open("wb") as binary:
        with gzip.GzipFile(filename="", mode="wb", fileobj=binary, mtime=0) as compressed:
            with io.TextIOWrapper(compressed, encoding="utf-8", newline="") as text:
                writer = csv.DictWriter(text, fieldnames=fields, lineterminator="\n")
                writer.writeheader()
                for spec in contract["tables"]:
                    table = spec["table"]
                    rows, columns, samples, values = load_table(paths[f"PXD057588/table{table}.xlsx"], spec)
                    totals = values.sum(axis=0)
                    scaled = values / totals[np.newaxis, :]
                    paper_values = scaled if spec["policy"] == "total_signal" else values
                    wrong_values = values if spec["policy"] == "total_signal" else scaled
                    paper_errors, wrong_errors, reported_values = [], [], []
                    parents = [row[column_index_from_string(spec["parent_column"]) - 1] for row in rows[5:]]
                    assert all(isinstance(parent, str) and parent for parent in parents)
                    counts = collections.Counter(parents)
                    parent_counts[table] = counts
                    support.append({
                        "table": table, "n_peptide_rows": len(parents), "n_parent_labels": len(counts),
                        "n_singleton_parent_labels": sum(n == 1 for n in counts.values()),
                        "n_multiple_peptide_parent_labels": sum(n >= 2 for n in counts.values()),
                        "n_peptide_rows_in_multiple_peptide_groups": sum(n for n in counts.values() if n >= 2),
                        "maximum_peptides_per_parent_label": max(counts.values()),
                        "parent_size_histogram": dict(sorted(collections.Counter(counts.values()).items())),
                    })
                    for contrast in spec["contrasts"]:
                        sample = contrast["sample"]
                        idx = samples[sample]
                        controls = [samples[name] for name in contrast["control_samples"]]
                        fc_column = column_index_from_string(contrast["fc_column"]) - 1
                        assert rows[4][fc_column] == contrast["expected_fc_header"]
                        reported = np.array([row[fc_column] for row in rows[5:]], dtype=np.float64)
                        assert np.isfinite(reported).all()
                        calculated = np.log2(paper_values[:, idx] / paper_values[:, controls].mean(axis=1))
                        wrong = np.log2(wrong_values[:, idx] / wrong_values[:, controls].mean(axis=1))
                        error, wrong_error = calculated - reported, wrong - reported
                        paper_errors.append(error)
                        wrong_errors.append(wrong_error)
                        reported_values.append(reported)
                        channels.append({"table": table, "sample": sample, "paper_policy": spec["policy"],
                                         "paper": errors(error, contract["absolute_tolerance_log2"]),
                                         "wrong": errors(wrong_error, contract["absolute_tolerance_log2"])})
                        for row_index, source_row in enumerate(rows[5:], start=6):
                            offset = row_index - 6
                            writer.writerow({
                                "table": table, "source_row": row_index, "sequence": source_row[0],
                                "parent_label": parents[offset], "sample": sample,
                                "intensity_cell": f"{get_column_letter(columns[idx] + 1)}{row_index}",
                                "fold_change_cell": f"{contrast['fc_column']}{row_index}",
                                "control_columns": ";".join(get_column_letter(columns[i] + 1) for i in controls),
                                "reported_log2fc": reported[offset], "paper_policy_log2fc": calculated[offset],
                                "wrong_policy_log2fc": wrong[offset], "paper_policy_error_log2": error[offset],
                                "wrong_policy_error_log2": wrong_error[offset],
                            })
                    reported = np.concatenate(reported_values)
                    summaries.append({
                        "table": table, "paper_policy": spec["policy"],
                        "paper": errors(np.concatenate(paper_errors), contract["absolute_tolerance_log2"]),
                        "wrong": errors(np.concatenate(wrong_errors), contract["absolute_tolerance_log2"]),
                        "reported_log2fc_mean": float(reported.mean()),
                        "reported_log2fc_std_population": float(reported.std()),
                        "reported_log2fc_minimum": float(reported.min()),
                        "reported_log2fc_maximum": float(reported.max()),
                    })
    assert sum(x["paper"]["n_cells"] for x in summaries) == contract["expected_n_fold_change_cells"]
    shared = set(parent_counts[9]) & set(parent_counts[11])
    shared_multiple = {p for p in shared if parent_counts[9][p] >= 2 and parent_counts[11][p] >= 2}
    overlap = {
        "n_shared_parent_labels": len(shared),
        "n_elution_rows_with_shared_parent_label": sum(parent_counts[9][p] for p in shared),
        "n_shared_labels_with_multiple_ligands_and_multiple_digest_peptides": len(shared_multiple),
        "n_elution_rows_in_those_groups": sum(parent_counts[9][p] for p in shared_multiple),
        "interpretation": "Literal source-label support only; not verified specimen or sequence-reference joins.",
    }
    save(results / "summary.json", summaries)
    save(results / "channel_reconciliation.json", channels)
    save(results / "parent_support.json", {"tables": support, "table9_table11_overlap": overlap})
    save(results / "runtime.json", {"seconds": time.perf_counter() - start,
         "python": platform.python_version(), "numpy": np.__version__, "openpyxl": openpyxl.__version__,
         "platform": platform.platform(), "gpu": None, "training": False})
    save(results / "output_hashes.json", {p.name: {"sha256": hashlib.sha256(p.read_bytes()).hexdigest(),
         "size_bytes": p.stat().st_size} for p in sorted(results.iterdir()) if p.name != "output_hashes.json"})
    print(json.dumps({"n_cells": contract["expected_n_fold_change_cells"],
                      "n_paper_policy_mismatches": sum(s["paper"]["n_mismatches"] for s in summaries),
                      "n_wrong_policy_mismatches": sum(s["wrong"]["n_mismatches"] for s in summaries),
                      "parent_overlap": overlap}))


if __name__ == "__main__":
    main()
