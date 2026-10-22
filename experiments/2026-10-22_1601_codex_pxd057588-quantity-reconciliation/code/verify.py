"""Independent scalar replay of original cells and every comparison-dump row."""

from collections import Counter
import csv
import gzip
import hashlib
import json
import math
from pathlib import Path
import time

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from reconcile import FAMILY, inputs


def main():
    start = time.perf_counter()
    contract = json.loads((FAMILY / "reproduce/contract.json").read_text())
    launch = json.loads((FAMILY / "reproduce/launch.json").read_text())
    assert hashlib.sha256((FAMILY / "reproduce/contract.json").read_bytes()).hexdigest() == launch["contract_sha256"]
    assert hashlib.sha256((FAMILY / "reproduce/source/reconcile.py").read_bytes()).hexdigest() == launch["source_sha256"]
    paths = inputs(contract)  # Only input retrieval is shared, no calculation code.
    results = FAMILY / "results"
    hashes = json.loads((results / "output_hashes.json").read_text())
    for name, receipt in hashes.items():
        path = results / name
        assert path.stat().st_size == receipt["size_bytes"]
        assert hashlib.sha256(path.read_bytes()).hexdigest() == receipt["sha256"]
    summary = {s["table"]: s for s in json.loads((results / "summary.json").read_text())}
    support = json.loads((results / "parent_support.json").read_text())
    groups = {}
    checked = 0
    maximum_replay_error = 0.0
    with gzip.open(results / "cell_comparisons.csv.gz", "rt", newline="") as handle:
        dump = csv.DictReader(handle)
        for spec in contract["tables"]:
            table = spec["table"]
            workbook = openpyxl.load_workbook(paths[f"PXD057588/table{table}.xlsx"], read_only=True, data_only=False)
            rows = list(workbook.active.values)
            workbook.close()
            assert len(rows) == spec["n_rows"] + 5
            columns = {}
            for i in range(1, len(rows[3])):
                if rows[3][i] is None:
                    break
                assert rows[4][i] not in columns
                columns[rows[4][i]] = i
            totals = {i: math.fsum(float(row[i]) for row in rows[5:]) for i in columns.values()}
            parent_index = column_index_from_string(spec["parent_column"]) - 1
            groups[table] = Counter(row[parent_index] for row in rows[5:])
            source_counts = groups[table]
            expected_support = next(s for s in support["tables"] if s["table"] == table)
            assert expected_support == {
                "table": table, "n_peptide_rows": len(rows) - 5,
                "n_parent_labels": len(source_counts),
                "n_singleton_parent_labels": sum(v == 1 for v in source_counts.values()),
                "n_multiple_peptide_parent_labels": sum(v > 1 for v in source_counts.values()),
                "n_peptide_rows_in_multiple_peptide_groups": sum(v for v in source_counts.values() if v > 1),
                "maximum_peptides_per_parent_label": max(source_counts.values()),
                "parent_size_histogram": {str(k): v for k, v in Counter(source_counts.values()).items()},
            }
            errors = {"paper": [], "wrong": []}
            for contrast in spec["contrasts"]:
                treated = columns[contrast["sample"]]
                controls = [columns[c] for c in contrast["control_samples"]]
                fc = column_index_from_string(contrast["fc_column"]) - 1
                assert rows[4][fc] == contrast["expected_fc_header"]
                for row_number, row in enumerate(rows[5:], start=6):
                    output = next(dump)
                    expected = {
                        "table": str(table), "source_row": str(row_number),
                        "sequence": row[0], "parent_label": row[parent_index],
                        "sample": contrast["sample"],
                        "intensity_cell": f"{get_column_letter(treated + 1)}{row_number}",
                        "fold_change_cell": f"{contrast['fc_column']}{row_number}",
                        "control_columns": ";".join(get_column_letter(c + 1) for c in controls),
                    }
                    assert all(output[k] == v for k, v in expected.items())
                    reported = float(row[fc])
                    raw = math.log2(float(row[treated]) / (math.fsum(float(row[c]) for c in controls) / len(controls)))
                    scaled = math.log2((float(row[treated]) / totals[treated]) / (
                        math.fsum(float(row[c]) / totals[c] for c in controls) / len(controls)))
                    paper, wrong = (scaled, raw) if table == 11 else (raw, scaled)
                    numbers = {
                        "reported_log2fc": reported,
                        "paper_policy_log2fc": paper, "wrong_policy_log2fc": wrong,
                        "paper_policy_error_log2": paper - reported,
                        "wrong_policy_error_log2": wrong - reported,
                    }
                    for key, value in numbers.items():
                        delta = abs(float(output[key]) - value)
                        assert delta < 1e-12, (table, row_number, key, delta)
                        maximum_replay_error = max(maximum_replay_error, delta)
                    errors["paper"].append(paper - reported)
                    errors["wrong"].append(wrong - reported)
                    checked += 1
            for name, values in errors.items():
                expected = summary[table][name]
                assert expected["n_cells"] == len(values)
                assert expected["n_mismatches"] == sum(abs(v) > contract["absolute_tolerance_log2"] for v in values)
                for key, value in {
                    "maximum_absolute_error_log2": max(abs(v) for v in values),
                    "mean_absolute_error_log2": math.fsum(abs(v) for v in values) / len(values),
                    "rmse_log2": math.sqrt(math.fsum(v * v for v in values) / len(values)),
                }.items():
                    assert abs(expected[key] - value) < 1e-12
        assert next(dump, None) is None
    assert checked == contract["expected_n_fold_change_cells"]
    shared = groups[9].keys() & groups[11].keys()
    multiple = {p for p in shared if groups[9][p] > 1 and groups[11][p] > 1}
    overlap = support["table9_table11_overlap"]
    assert overlap["n_shared_parent_labels"] == len(shared)
    assert overlap["n_elution_rows_with_shared_parent_label"] == sum(groups[9][p] for p in shared)
    assert overlap["n_shared_labels_with_multiple_ligands_and_multiple_digest_peptides"] == len(multiple)
    assert overlap["n_elution_rows_in_those_groups"] == sum(groups[9][p] for p in multiple)
    receipt = {
        "passed": True, "n_comparisons": checked, "n_original_workbooks": 6,
        "n_hashed_outputs": len(hashes), "n_parent_support_tables": len(groups),
        "maximum_scalar_vs_vector_delta_log2": maximum_replay_error,
        "seconds": time.perf_counter() - start,
        "verifier_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    }
    (FAMILY / "verification").mkdir(exist_ok=True)
    (FAMILY / "verification/independent_verification.json").write_text(json.dumps(receipt, indent=2) + "\n")
    print(json.dumps(receipt))


if __name__ == "__main__":
    main()
