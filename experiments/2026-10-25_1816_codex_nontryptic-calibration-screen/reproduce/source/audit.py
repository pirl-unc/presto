"""Source-supported dose qualification, with unresolved chemical identities retained."""

import argparse
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import openpyxl


def audit(root, data):
    book = openpyxl.load_workbook(data / "pnas.2111173118.sd06.xlsx", data_only=True)
    amap = {
        book["DatasetS6A"].cell(r, 1).value: (
            book["DatasetS6A"].cell(r, 3).value,
            f"C{r}",
        )
        for r in range(5, 15)
    }
    identities = []
    sheet = book["DatasetS6B"]
    for row in range(1, sheet.max_row + 1):
        isotope = sheet.cell(row, 1).value
        if isotope not in ("1H", "2H", "3H", "4H"):
            continue
        form = sheet.cell(row, 2).value
        sites = re.findall(r"\(([A-Z])\)", form)
        missing = sorted(set(sites) - set(amap))
        expected = None if missing else sum(amap[a][0] for a in sites)
        reported = sheet.cell(row, 4).value
        identities.append(
            {
                "isotopologue": isotope,
                "annotated_sequence": form,
                "sequence": form.replace("(", "").replace(")", ""),
                "form_cell": f"B{row}",
                "reported_nominal_shift_cell": f"D{row}",
                "reported_nominal_shift": reported,
                "sum_S6A_shift": expected,
                "S6A_mass_cells": [amap[a][1] for a in sites if a in amap],
                "status": "source_assignment_absent"
                if missing
                else "source_conflict"
                if abs(expected - reported) > 0.15
                else "nominal_shift_consistent",
                "canonical_chemical_identity_asserted": False,
            }
        )
    assert len(identities) == 72
    records = json.loads((root / "results/dose-signal-cells.json").read_text())
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    # An independent reader checks values by original cell, without block-filling.
    with ZipFile(data / "pnas.2111173118.sd03.xlsx") as z:
        maps = {}
        for name, member in (
            ("DatasetS3A", "xl/worksheets/sheet1.xml"),
            ("DatasetS3B", "xl/worksheets/sheet2.xml"),
        ):
            maps[name] = {
                c.attrib["r"]: c.findtext("s:v", namespaces=ns)
                for c in ET.fromstring(z.read(member)).findall(
                    ".//s:sheetData/s:row/s:c", ns
                )
            }
        for record in records:
            original = maps[record["sheet"]]
            raw = original.get(record["source_cell"])
            assert record["raw_signal"] == (None if raw is None else float(raw))
            assert record["added_fmol"] == float(original[record["dose_cell"]])
    out = {
        "n_original_response_cells_verified": len(records),
        "n_isotopologue_records": len(identities),
        "n_conflicting_isotopologues": sum(
            x["status"] == "source_conflict" for x in identities
        ),
        "identities": identities,
        "omitted_dose": {
            "sheet": "DatasetS3A",
            "label_cell": "A16",
            "sequence": "GLFDQHFRL",
            "fmol": 0.1,
            "source_row": None,
            "reason": "not_reported_no_zero_imputation",
        },
        "training_admitted": False,
    }
    with (root / "results/semantic-independent-audit.json").open("x") as f:
        json.dump(out, f, indent=2, allow_nan=False)
        f.write("\n")
    print(json.dumps({k: v for k, v in out.items() if k != "identities"}))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--requests", type=Path, required=True)
    parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    data = Path(
        json.loads((root / "reproduce/launch.json").read_text())["data_directory"]
    )
    audit(root, data)


if __name__ == "__main__":
    main()
