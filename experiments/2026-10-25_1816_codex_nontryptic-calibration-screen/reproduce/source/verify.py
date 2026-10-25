"""Independent literal-cell inventory; no calibration fits or biological scores."""

import argparse
import hashlib
import json
import posixpath
import re
import subprocess
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pdfplumber
from pypdf import PdfReader


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(node):
    return "".join(node.itertext()) if node is not None else ""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--requests", type=Path, required=True)
    request = json.loads(parser.parse_args().requests.read_text())
    root = Path(__file__).resolve().parents[1]
    config = json.loads((root / "reproduce/launch.json").read_text())
    data = Path(config["data_directory"])
    manifest = json.loads((data / "PMC8449407.1.json").read_text())
    assert (
        manifest["doi"] == "10.1073/pnas.2111173118"
        and manifest["pmcid"] == "PMC8449407"
    )
    assert (
        hashlib.md5((data / "PMC8449407.1.xml").read_bytes()).hexdigest()
        == manifest["xml_url"].split("md5=")[1]
    )
    xml = ET.parse(data / "PMC8449407.1.xml").getroot()
    assert [
        n.text
        for n in xml.findall(".//article-id")
        if n.attrib.get("pub-id-type") == "doi"
    ] == [manifest["doi"]]
    summaries = []
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    for filename in request["files"]:
        path = data / filename
        url = [u for u in manifest["media_urls"] if "/" + filename + "?" in u]
        assert len(url) == 1
        assert hashlib.md5(path.read_bytes()).hexdigest() == url[0].split("md5=")[1]
        receipt = json.loads((data / (filename + ".receipt.json")).read_text())
        assert (
            receipt["status"] == "retrieved_validated_payload"
            and sha(path) == receipt["sha256"]
        )
        if not filename.endswith(".xlsx"):
            continue
        book = openpyxl.load_workbook(path, data_only=True, read_only=False)
        formula_book = openpyxl.load_workbook(path, data_only=False, read_only=False)
        with ZipFile(path) as archive:
            strings = []
            if "xl/sharedStrings.xml" in archive.namelist():
                strings = [
                    "".join(n.itertext())
                    for n in ET.fromstring(
                        archive.read("xl/sharedStrings.xml")
                    ).findall("s:si", ns)
                ]
            relations = {
                n.attrib["Id"]: n.attrib["Target"]
                for n in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            }
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            for sheet in workbook.findall("s:sheets/s:sheet", ns):
                name = sheet.attrib["name"]
                relation = sheet.attrib[
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                ]
                target = relations[relation]
                member = (
                    posixpath.normpath(posixpath.join("xl", target))
                    if not target.startswith("/")
                    else target.lstrip("/")
                )
                worksheet = ET.fromstring(archive.read(member))
                records = []
                formulas = []
                types = Counter()
                for cell in worksheet.findall(".//s:sheetData/s:row/s:c", ns):
                    coord = cell.attrib["r"]
                    kind = cell.attrib.get("t", "n")
                    value = cell.findtext("s:v", namespaces=ns)
                    if kind == "s":
                        decoded = strings[int(value)] if value is not None else None
                    elif kind == "inlineStr":
                        decoded = "".join(
                            n.text or "" for n in cell.findall(".//s:t", ns)
                        )
                    elif kind == "b":
                        decoded = value == "1" if value is not None else None
                    elif kind in ("str", "e"):
                        decoded = value
                    elif value is None:
                        decoded = None
                    else:
                        decoded = float(value)
                    actual = book[name][coord].value
                    if isinstance(actual, datetime):
                        assert kind == "n" and decoded is not None
                        offset = (
                            decoded
                            if book.epoch.year == 1904
                            else decoded - (1 if decoded >= 60 else 0)
                        )
                        start = (
                            book.epoch
                            if book.epoch.year == 1904
                            else datetime(1899, 12, 31)  # noqa: DTZ001 -- Excel serials have no timezone
                        )
                        assert actual == start + timedelta(days=offset)
                    else:
                        assert actual == decoded, (
                            filename,
                            name,
                            coord,
                            actual,
                            decoded,
                            kind,
                        )
                    formula = cell.find("s:f", ns)
                    if formula is not None:
                        formulas.append(
                            {
                                "cell": coord,
                                "expression": formula.text,
                                "cached_value": value,
                            }
                        )
                        assert formula_book[name][coord].data_type == "f"
                    types[kind] += 1
                    records.append(
                        {
                            "cell": coord,
                            "type": kind,
                            "style": cell.attrib.get("s"),
                            "raw_value": value,
                            "value": str(actual)
                            if isinstance(actual, datetime)
                            else actual,
                        }
                    )
                preview = []
                for row in book[name].iter_rows():
                    values = [
                        {
                            "cell": c.coordinate,
                            "value": str(c.value)
                            if isinstance(c.value, datetime)
                            else c.value,
                        }
                        for c in row
                        if c.value is not None
                    ]
                    if values:
                        preview.append(values)
                    if len(preview) >= 14:
                        break
                summaries.append(
                    {
                        "file": filename,
                        "sheet": name,
                        "dimensions": book[name].calculate_dimension(),
                        "n_stored_cells_replayed": len(records),
                        "cell_types": dict(types),
                        "formulas": formulas,
                        "literal_cells_sha256": hashlib.sha256(
                            json.dumps(
                                records, sort_keys=True, ensure_ascii=False
                            ).encode()
                        ).hexdigest(),
                        "first_nonempty_rows": preview,
                    }
                )
        book.close()
        formula_book.close()
    pdf = data / request["pdf"]
    doc = pdfplumber.open(pdf)
    reader = PdfReader(pdf)
    assert len(doc.pages) == len(reader.pages)
    extraction = data / "appendix_inspection"
    extraction.mkdir(exist_ok=False)
    pages = []
    for i, page in enumerate(doc.pages):
        left = page.extract_text()
        right = reader.pages[i].extract_text()
        (extraction / f"page-{i + 1}-pdfplumber.txt").write_text(left)
        (extraction / f"page-{i + 1}-pypdf.txt").write_text(right)
        normal = lambda t: Counter(re.findall(r"[A-Za-z0-9]+", t.lower()))
        a, b = normal(left), normal(right)
        pages.append(
            {
                "page": i + 1,
                "pdfplumber_word_count": sum(a.values()),
                "pypdf_word_count": sum(b.values()),
                "shared_word_multiset_count": sum((a & b).values()),
                "contains_calibration_terms": [
                    w
                    for w in (
                        "purity",
                        "ELISA",
                        "isotopologue",
                        "concentration",
                        "Skyline",
                        "Dataset S3",
                        "fmol",
                    )
                    if w.lower() in left.lower() and w.lower() in right.lower()
                ],
            }
        )
    doc.close()
    subprocess.run(
        [
            "/opt/homebrew/bin/pdftoppm",
            "-r",
            "100",
            "-png",
            str(pdf),
            str(extraction / "page"),
        ],
        check=True,
    )
    result = {
        "status": "literal_inventory_verified_semantic_admission_pending",
        "doi": manifest["doi"],
        "accession": "PXD024917",
        "n_workbooks": 3,
        "n_sheets": len(summaries),
        "n_source_cells_replayed": sum(s["n_stored_cells_replayed"] for s in summaries),
        "workbooks": summaries,
        "pdf_pages": pages,
        "pdf_extractions_and_renders": str(extraction),
        "all_official_md5s_verified": True,
        "model_fitting_or_scoring": False,
    }
    with (root / "results/source-inventory.json").open("x") as f:
        json.dump(result, f, indent=2, allow_nan=False)
        f.write("\n")
    print(json.dumps({k: v for k, v in result.items() if k != "workbooks"}))


if __name__ == "__main__":
    main()
