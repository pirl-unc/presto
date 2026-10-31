"""Source-supported dose qualification, with unresolved chemical identities retained."""

import argparse
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

import openpyxl


def semantic(root, data):
    import math
    import re

    panel_book = openpyxl.load_workbook(
        data / "pnas.2111173118.sd01.xlsx", data_only=True
    )
    panel = {panel_book.active.cell(row, 2).value for row in range(4, 22)}
    assert len(panel) == 18
    source = openpyxl.load_workbook(data / "pnas.2111173118.sd03.xlsx", data_only=True)
    filemap = openpyxl.load_workbook(data / "pnas.2111173118.sd07.xlsx", data_only=True)
    file_rows = [
        list(row)
        for row in filemap.active.iter_rows(min_row=2, max_col=3, values_only=True)
        if row[0] is not None
    ]
    records = []
    groups = []
    forms = {}
    for sheet_name, expected_doses, spike_stage in (
        (
            "DatasetS3A",
            [0.1, 1.0, 10.0, 100.0],
            "free_peptide_after_HLA_peptide_purification",
        ),
        (
            "DatasetS3B",
            [0.1, 1.0, 10.0, 100.0, 1000.0],
            "hipMHC_before_immunoprecipitation",
        ),
    ):
        sheet = source[sheet_name]
        assert str(sheet.merged_cells) == "C3:E3"
        current = None
        current_rows = []
        for row in range(4, sheet.max_row + 1):
            label = sheet.cell(row, 1).value
            dose = sheet.cell(row, 2).value
            if dose is None:
                assert all(sheet.cell(row, col).value is None for col in range(3, 6))
                continue
            if label is not None:
                if current is not None:
                    assert [r[1] for r in current_rows] == (
                        [1.0, 10.0, 100.0]
                        if sheet_name == "DatasetS3A" and current == "GLFDQHFRL[+7]"
                        else expected_doses
                    )
                    groups.append((sheet_name, current, current_rows))
                current = str(label)
                current_rows = []
            assert current is not None and isinstance(dose, (float, int))
            sequence = re.sub(r"\[\+[-0-9.]+\]", "", current)
            assert sequence in panel, (sheet_name, current, sequence)
            current_rows.append((row, float(dose)))
            forms[(sheet_name, current)] = sequence
            for col in range(3, 6):
                cell = sheet.cell(row, col)
                value = cell.value
                assert (
                    value is None
                    or isinstance(value, (int, float))
                    and math.isfinite(value)
                    and value >= 0
                )
                records.append(
                    {
                        "sheet": sheet_name,
                        "source_cell": cell.coordinate,
                        "source_peptide_block": current,
                        "sequence": sequence,
                        "dose_cell": sheet.cell(row, 2).coordinate,
                        "added_fmol": float(dose),
                        "raw_signal": value,
                        "signal_status": "not_reported"
                        if value is None
                        else "reported_zero"
                        if value == 0
                        else "reported_positive",
                        "replicate_column_ordinal": col - 2,
                        "acquisition_id": None,
                        "acquisition_link_status": "unlabeled_source_column_no_exact_raw_file_join",
                        "spike_stage": spike_stage,
                        "isotopologue": "1H" if sheet_name.endswith("A") else "2H",
                    }
                )
        assert [r[1] for r in current_rows] == (
            [1.0, 10.0, 100.0]
            if sheet_name == "DatasetS3A" and current == "GLFDQHFRL[+7]"
            else expected_doses
        )
        groups.append((sheet_name, current, current_rows))
    summary = {}
    for sheet_name in ("DatasetS3A", "DatasetS3B"):
        chosen = [r for r in records if r["sheet"] == sheet_name]
        sequences = {r["sequence"] for r in chosen}
        n_supported = sum(
            len(
                {
                    r["added_fmol"]
                    for r in chosen
                    if r["sequence"] == seq
                    and r["signal_status"] == "reported_positive"
                }
            )
            >= 3
            for seq in sequences
        )
        summary[sheet_name] = {
            "n_sequences": len(sequences),
            "n_slots": len(chosen),
            "n_positive": sum(
                r["signal_status"] == "reported_positive" for r in chosen
            ),
            "n_zero": sum(r["signal_status"] == "reported_zero" for r in chosen),
            "n_missing": sum(r["signal_status"] == "not_reported" for r in chosen),
            "n_sequences_with_3_positive_doses": n_supported,
        }
    assert summary["DatasetS3A"]["n_sequences_with_3_positive_doses"] >= 12
    source_exports = [r for r in file_rows if str(r[1]).lower() == "search"]
    result = {
        "status": "nominal_amount_signal_support_qualified_training_metadata_unresolved",
        "cohort_support": summary,
        "n_panel_sequences": 18,
        "filemap_source_exports": source_exports,
        "scope": "selected HLA-A0201 tumor-antigen targeted MS2 response; not MS1 or intracellular free-peptide concentration",
        "metadata_gaps": [
            "unlabeled source response columns need exact acquisition mapping",
            "transition identities/aggregation/interference need deposited transition exports",
            "18 selected targets do not establish broad sequence or protocol transfer",
            "S6A valine isotope assignment conflicts with S3/S6B nominal mass shift",
        ],
        "training_admitted": False,
        "shared_hitlist_ingestion_required": True,
        "amount_assignment": {
            "free_peptide": "reported amino-acid analysis; HPLC purity >97%; isotopologues balanced by LC/MS",
            "hipMHC": "reported stable-complex ELISA assignment after exchange",
            "source": "original supplementary Methods page2; not independently certified values",
        },
        "forward_fill_scope": "peptide label applies to its explicitly delimited table block with the source-reported dose grid; GLFDQHFRL[+7] has no 0.1 fmol source row; original blank cells unchanged",
        "measurement_column_semantics": "ordinal worksheet slot only, not asserted technical injection identity",
    }
    with (root / "results/dose-signal-cells.json").open("x") as f:
        json.dump(records, f, indent=2, allow_nan=False)
        f.write("\n")
    with (root / "results/semantic-qualification.json").open("x") as f:
        json.dump(result, f, indent=2, allow_nan=False)
        f.write("\n")
    print(json.dumps(result))


def audit(root, data):
    book = openpyxl.load_workbook(data / "pnas.2111173118.sd06.xlsx", data_only=True)
    amap = {
        book["DatasetS6A"].cell(r, 1).value: (
            book["DatasetS6A"].cell(r, 3).value,
            f"C{r}",
        )
        for r in range(5, 15)
    }
    identities = []
    sheet = book["DatasetS6B"]
    for row in range(1, sheet.max_row + 1):
        isotope = sheet.cell(row, 1).value
        if isotope not in ("1H", "2H", "3H", "4H"):
            continue
        form = sheet.cell(row, 2).value
        sites = re.findall(r"\(([A-Z])\)", form)
        expected = sum(amap[a][0] for a in sites)
        reported = sheet.cell(row, 4).value
        identities.append(
            {
                "isotopologue": isotope,
                "annotated_sequence": form,
                "sequence": form.replace("(", "").replace(")", ""),
                "form_cell": f"B{row}",
                "reported_nominal_shift_cell": f"D{row}",
                "reported_nominal_shift": reported,
                "sum_S6A_shift": expected,
                "S6A_mass_cells": [amap[a][1] for a in sites],
                "status": "source_conflict"
                if abs(expected - reported) > 0.15
                else "nominal_shift_consistent",
                "canonical_chemical_identity_asserted": False,
            }
        )
    assert len(identities) == 72
    records = json.loads((root / "results/dose-signal-cells.json").read_text())
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    # An independent reader checks values by original cell, without block-filling.
    with ZipFile(data / "pnas.2111173118.sd03.xlsx") as z:
        maps = {}
        for name, member in (
            ("DatasetS3A", "xl/worksheets/sheet1.xml"),
            ("DatasetS3B", "xl/worksheets/sheet2.xml"),
        ):
            maps[name] = {
                c.attrib["r"]: c.findtext("s:v", namespaces=ns)
                for c in ET.fromstring(z.read(member)).findall(
                    ".//s:sheetData/s:row/s:c", ns
                )
            }
        for record in records:
            original = maps[record["sheet"]]
            raw = original.get(record["source_cell"])
            assert record["raw_signal"] == (None if raw is None else float(raw))
            assert record["added_fmol"] == float(original[record["dose_cell"]])
    out = {
        "n_original_response_cells_verified": len(records),
        "n_isotopologue_records": len(identities),
        "n_conflicting_isotopologues": sum(
            x["status"] == "source_conflict" for x in identities
        ),
        "identities": identities,
        "omitted_dose": {
            "sheet": "DatasetS3A",
            "label_cell": "A16",
            "sequence": "GLFDQHFRL",
            "fmol": 0.1,
            "source_row": None,
            "reason": "not_reported_no_zero_imputation",
        },
        "training_admitted": False,
    }
    with (root / "results/semantic-independent-audit.json").open("x") as f:
        json.dump(out, f, indent=2, allow_nan=False)
        f.write("\n")
    print(json.dumps({k: v for k, v in out.items() if k != "identities"}))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--requests", type=Path, required=True)
    parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    data = Path(
        json.loads((root / "reproduce/launch.json").read_text())["data_directory"]
    )
    semantic(root, data)
    audit(root, data)


if __name__ == "__main__":
    main()
