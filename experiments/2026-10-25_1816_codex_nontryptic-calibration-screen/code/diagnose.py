"""Read original dose blocks and metadata without assuming a uniform dose grid."""

import argparse
import json
from pathlib import Path

import openpyxl


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--requests", type=Path, required=True)
    parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    config = json.loads((root / "reproduce/launch.json").read_text())
    data = Path(config["data_directory"])
    source = openpyxl.load_workbook(data / "pnas.2111173118.sd03.xlsx", data_only=True)
    blocks = []
    for name in ("DatasetS3A", "DatasetS3B"):
        sheet = source[name]
        current = None
        for row in range(4, sheet.max_row + 1):
            label = sheet.cell(row, 1).value
            dose = sheet.cell(row, 2).value
            if label is not None:
                current = {
                    "sheet": name,
                    "label_cell": f"A{row}",
                    "label": label,
                    "rows": [],
                }
                blocks.append(current)
            if dose is not None:
                assert current is not None
                current["rows"].append(
                    {
                        "dose_cell": f"B{row}",
                        "dose": dose,
                        "missing_response_cells": [
                            sheet.cell(row, col).coordinate
                            for col in range(3, 6)
                            if sheet.cell(row, col).value is None
                        ],
                    }
                )
    metadata = {}
    for filename in ("pnas.2111173118.sd06.xlsx", "pnas.2111173118.sd07.xlsx"):
        book = openpyxl.load_workbook(data / filename, data_only=True)
        for sheet in book:
            metadata[filename + ":" + sheet.title] = [
                [
                    {"cell": c.coordinate, "value": c.value}
                    for c in row
                    if c.value is not None
                ]
                for row in sheet.iter_rows()
                if any(c.value is not None for c in row)
            ]
    result = {"blocks": blocks, "source_metadata": metadata, "model_fits": 0}
    with (root / "results/layout-diagnosis.json").open("x") as out:
        json.dump(result, out, indent=2, allow_nan=False)
        out.write("\n")
    for b in blocks:
        print(b["sheet"], b["label_cell"], b["label"], [r["dose"] for r in b["rows"]])


if __name__ == "__main__":
    main()
