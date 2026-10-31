"""Read-only workbook inventory and independent OOXML value replay; no cohort."""

import collections
import datetime
import gzip
import hashlib
import json
import posixpath
import re
import resource
import shutil
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
EXP = ROOT / "experiments/2026-10-25_0344_codex_proteasome-quantity-source"
RAW = ROOT / "data/proteasome_quantity_sources" / EXP.name
OUT = EXP / "results"
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
START = time.monotonic()


def check_resources():
    assert time.monotonic() - START < 1200, "Analysis exceeded 20 minutes"
    assert resource.getrusage(resource.RUSAGE_SELF).ru_maxrss < 3 * 1024**3
    assert shutil.disk_usage(RAW).free >= 1024**3 + 32 * 1024**2


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def xml_sheets(archive):
    shared = []
    if "xl/sharedStrings.xml" in archive.namelist():
        shared = [
            "".join(node.itertext()) for node in ET.fromstring(archive.read("xl/sharedStrings.xml"))
        ]
    relations = {
        node.attrib["Id"]: node.attrib["Target"]
        for node in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    }
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    for sheet in workbook.findall("s:sheets/s:sheet", NS):
        target = relations[sheet.attrib[RID]]
        target = target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target)
        tree = ET.fromstring(archive.read(posixpath.normpath(target)))
        cells = {}
        shared_formulas = {}
        for node in tree.findall("s:sheetData/s:row/s:c", NS):
            kind = node.attrib.get("t", "n")
            raw_value = node.findtext("s:v", default=None, namespaces=NS)
            formula = node.find("s:f", NS)
            if formula is not None and formula.attrib.get("t") == "shared" and formula.text:
                shared_formulas[formula.attrib["si"]] = (node.attrib["r"], formula.text)
            if kind == "s":
                value = shared[int(raw_value)] if raw_value is not None else None
            elif kind == "inlineStr":
                inline = node.find("s:is", NS)
                value = "".join(inline.itertext()) if inline is not None else None
            elif kind == "b":
                value = bool(int(raw_value)) if raw_value is not None else None
            elif kind == "n":
                value = float(raw_value) if raw_value is not None else None
            else:
                value = raw_value
            cells[node.attrib["r"]] = {
                "xml_type": kind,
                "raw_value": raw_value,
                "style": node.attrib.get("s"),
                "value": value,
                "formula": formula.text if formula is not None else None,
                "formula_attributes": dict(formula.attrib) if formula is not None else None,
            }
        for coord, item in cells.items():
            attrs = item["formula_attributes"]
            item["resolved_formula"] = item["formula"]
            if attrs is not None and attrs.get("t") == "shared" and item["formula"] is None:
                origin, expression = shared_formulas[attrs["si"]]
                origin_col, origin_row = re.fullmatch(r"([A-Z]+)([0-9]+)", origin).groups()
                col, row = re.fullmatch(r"([A-Z]+)([0-9]+)", coord).groups()
                assert col == origin_col, "This audit only qualifies vertical shared formulas"
                assert '"' not in expression and "'" not in expression
                delta = int(row) - int(origin_row)
                item["resolved_formula"] = re.sub(
                    r"(\$?[A-Z]+)(\$?)([0-9]+)",
                    lambda m, delta=delta: m[1] + m[2] + str(int(m[3]) + (0 if m[2] else delta)),
                    expression,
                )
        yield sheet.attrib["name"], cells


def main():
    OUT.mkdir(exist_ok=True)
    destination = OUT / "workbook-inventory.json"
    assert not destination.exists(), "Never overwrite an earlier audit"
    dump_path = RAW / "independent-original-cells-shared-formula.jsonl.gz"
    assert not dump_path.exists()
    check_resources()
    summary = {"files": [], "mismatches": [], "date_cells": [], "source_only": True}
    with gzip.open(dump_path, "xt", encoding="utf-8") as dump:
        for path in sorted((RAW / "complete-members").glob("*.xlsx")):
            cached = load_workbook(path, read_only=True, data_only=True)
            formulas = load_workbook(path, read_only=True, data_only=False)
            file_record = {"file": path.name, "sha256": sha(path), "sheets": []}
            with zipfile.ZipFile(path) as archive:
                assert archive.testzip() is None
                for name, cells in xml_sheets(archive):
                    check_resources()
                    values = {
                        c.coordinate: c.value
                        for row in cached[name]
                        for c in row
                        if c.value is not None
                    }
                    formula_values = {
                        c.coordinate: c.value
                        for row in formulas[name]
                        for c in row
                        if c.data_type == "f"
                    }
                    nonempty = 0
                    for coord, item in cells.items():
                        dump.write(
                            json.dumps({"file": path.name, "sheet": name, "cell": coord, **item})
                            + "\n"
                        )
                        value = item["value"]
                        actual = values.pop(coord, None)
                        nonempty += value is not None or item["formula_attributes"] is not None
                        if isinstance(actual, (datetime.datetime, datetime.date, datetime.time)):
                            summary["date_cells"].append(
                                [path.name, name, coord, str(actual), item]
                            )
                        elif actual != value:
                            summary["mismatches"].append(
                                [path.name, name, coord, "cached", actual, value]
                            )
                        if item["resolved_formula"] is not None:
                            observed = formula_values.pop(coord, None)
                            expected = "=" + item["resolved_formula"]
                            if observed != expected:
                                summary["mismatches"].append(
                                    [path.name, name, coord, "formula", str(observed), expected]
                                )
                    assert not values, (path.name, name, "unreplayed values", len(values))
                    assert not formula_values
                    counts = collections.Counter(
                        item["xml_type"] for item in cells.values() if item["value"] is not None
                    )
                    first_rows = [
                        {c.coordinate: c.value for c in row if isinstance(c.value, str) and c.data_type != "f"}
                        for row in formulas[name].iter_rows(min_row=1, max_row=8)
                    ]
                    record = {
                        "sheet": name,
                        "xml_cells": len(cells),
                        "nonempty_cells": nonempty,
                        "reported_rows": cached[name].max_row,
                        "reported_columns": cached[name].max_column,
                        "value_types": dict(counts),
                        "formulas": sum(
                            x["formula_attributes"] is not None for x in cells.values()
                        ),
                        "formulas_without_cache": sum(
                            x["formula_attributes"] is not None and x["value"] is None
                            for x in cells.values()
                        ),
                        "error_values": dict(
                            collections.Counter(
                                x["value"] for x in cells.values() if x["xml_type"] == "e"
                            )
                        ),
                        "first_rows_text_only": first_rows,
                    }
                    file_record["sheets"].append(record)
            cached.close()
            formulas.close()
            assert sha(path) == file_record["sha256"], "Original workbook changed"
            summary["files"].append(file_record)
            print(path.name, "sheets", len(file_record["sheets"]), flush=True)
    check_resources()
    summary.update(
        {
            "seconds": time.monotonic() - START,
            "peak_rss_bytes": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss,
            "independent_original_cells": str(dump_path),
            "independent_original_cells_sha256": sha(dump_path),
            "all_cells_match": not summary["mismatches"] and not summary["date_cells"],
        }
    )
    destination.write_text(json.dumps(summary, indent=2, default=str) + "\n")
    print(json.dumps({k: v for k, v in summary.items() if k != "files"}, default=str))
    assert summary["all_cells_match"], "Inspect source-format differences before admission"


if __name__ == "__main__":
    main()
