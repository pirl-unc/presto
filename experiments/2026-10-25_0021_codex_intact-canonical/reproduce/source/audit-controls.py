"""Independent original-workbook replay through public EvidenceStore queries."""

import collections
import hashlib
import importlib.util
import json
import math
import re
import resource
import time
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from hitlist.evidence_store import EvidenceStore

ROOT = Path(__file__).resolve().parents[1]
CONFIG = json.loads((ROOT / "reproduce/launch.json").read_text())
RAW = Path(CONFIG["original_assets"])
UPSTREAM = Path(CONFIG["upstream_audit"])
START = time.monotonic()
assert json.loads((ROOT / "results/build-resource.json").read_text())["passed"]
store = EvidenceStore.open(CONFIG["package"])


def query(name, **kwargs):
    return store.query("intact_" + name, limit=None, **kwargs).table.to_pylist()


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def raw_text(value):
    return None if value is None else str(value)


def equal_number(actual, expected):
    assert actual == (None if expected is None else float(expected)), (actual, expected)


spec = importlib.util.spec_from_file_location("original_xml_audit", UPSTREAM / "code/inventory.py")
original = importlib.util.module_from_spec(spec)
spec.loader.exec_module(original)
previous = json.loads((UPSTREAM / "results/quantity-audit-layout-qualified.json").read_text())
inventory = json.loads((UPSTREAM / "results/workbook-inventory.json").read_text())
assert inventory["all_cells_match"]
records = {
    x["source_record_id"]: (x["source_file"], x["sheet_name"], x["source_row_index"])
    for x in query("source_records", select=["source_record_id", "source_file", "sheet_name",
                                            "source_row_index"])
}
views = {x["view_id"]: x for x in query("source_views")}
acquisitions = {x["acquisition_id"]: x for x in query("acquisitions")}
quantities = {x["measurement_id"]: x for x in query("quantities")}
quantities_by_file = collections.defaultdict(list)
for q in quantities.values():
    quantities_by_file[q["source_file"]].append(q)
counts = collections.Counter()
workbooks = []

for book in inventory["files"]:
    filename = book["file"]
    assert digest(RAW / filename) == book["sha256"]
    actual = {(r["sheet_name"], r["source_cell"]): r
              for r in query("source_cells", where={"source_file": filename})}
    expected = {}
    with zipfile.ZipFile(RAW / filename) as archive:
        for sheet, cells in original.xml_sheets(archive):
            for coordinate, cell in cells.items():
                row = actual.pop((sheet, coordinate))
                col, number = coordinate_from_string(coordinate)
                assert row["source_column_index"] == column_index_from_string(col)
                assert row["source_row_index"] == number
                assert row["cell_type_raw"] == cell["xml_type"]
                assert row["cell_style_raw"] == cell["style"]
                assert row["cell_value_raw"] == cell["raw_value"]
                assert row["cell_formula_raw"] == cell["formula"]
                attrs = row["formula_attributes_json"]
                assert (None if attrs is None else json.loads(attrs)) == cell["formula_attributes"]
                resolved = cell["value"] if cell["xml_type"] in {"s", "inlineStr"} else cell["raw_value"]
                assert row["cell_value_resolved"] == resolved
                assert records[row["source_record_id"]] == (filename, sheet, number)
                expected[sheet, coordinate] = cell
                counts["n_source_cells"] += 1
                counts["n_populated_source_cells"] += cell["value"] is not None or cell["formula_attributes"] is not None
                counts["n_formulas"] += cell["formula_attributes"] is not None
    assert not actual, (filename, "extra source cells")
    for q in quantities_by_file[filename]:
        cell = expected.get((q["sheet_name"], q["source_cell"]), {})
        source_value = cell.get("value")
        resolved = source_value if cell.get("xml_type") in {"s", "inlineStr"} else cell.get("raw_value")
        assert q["quantitative_raw"] == resolved
        try:
            value = float(source_value) if source_value is not None else None
            value = value if value is None or math.isfinite(value) else None
        except (TypeError, ValueError):
            value = None
        if cell.get("xml_type") in {"e", "b", "d"}:
            value = None
        equal_number(q["quantitative_value"], value)
        if cell.get("formula_attributes") is not None:
            assert q["value_status"] == ("cached_formula_value" if value is not None else "formula_without_numeric_cache")
        elif cell.get("xml_type") in {"e", "b", "d"}:
            assert q["value_status"] == "source_error_or_nonquantity_type"
        elif value is not None:
            assert q["value_status"] == ("positive" if value > 0 else "zero" if value == 0 else "negative")
        assert q["quantity_scale"] == "source_linear_signal_not_molar"
        counts["n_quantities"] += 1
    workbooks.append({"source_file": filename, "n_source_cells": len(expected),
                      "n_quantities": len(quantities_by_file[filename]), "sha256": book["sha256"]})
assert counts["n_populated_source_cells"] == 325718 and counts["n_formulas"] == 2991
assert counts["n_quantities"] == len(quantities) == 44457

# Independent openpyxl body interpretation, using the already qualified source exceptions.
mass_rows = {records[x["source_record_id"]]: x for x in query("input_masses")}
assignment_rows = {records[x["source_record_id"]]: x for x in query("assignments")}
assignment_ids = {x["assignment_id"]: x for x in assignment_rows.values()}
input_ids = {x["input_mass_id"]: x for x in mass_rows.values()}
expected_quantities = set()
for filename in ("F1_SF2_SF3_SF4_SF5_SearchResults.xlsx", "F2_SF6_SearchResults.xlsx", "F3_SF7_SearchResults.xlsx"):
    workbook = load_workbook(RAW / filename, read_only=True, data_only=True)
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        cols = {h: i for i, h in enumerate(header) if h is not None}
        shift = filename.startswith("F2_") and sheet.title == "20S_AlphaSyn_40min_R1"
        no_ids = filename.startswith("F3_") and sheet.title == "Brain_20S_aSyn_0min_2"
        cols = {h: i + int(shift and i >= 7) for h, i in cols.items()}
        local_assignments, local_inputs = [], {}
        for index, row in enumerate(rows, 2):
            key = filename, sheet.title, index
            if row[0] is not None:
                m = mass_rows.pop(key)
                native = None if no_ids else row[0]
                assert m["native_mass_id_raw"] == raw_text(native)
                equal_number(m["mass_y"], row[0 if no_ids else 1])
                equal_number(m["mass_adjusted"], row[1 if no_ids else 2])
                q = quantities[m["measurement_id"]]
                equal_number(q["quantitative_value"], row[2 if no_ids else 3])
                expected_quantities.add(q["measurement_id"])
                if native is not None:
                    local_inputs[str(int(native))] = m
                counts["n_input_mass_rows"] += 1
            if "Sequence" not in cols or row[cols["Sequence"]] is None:
                continue
            a = assignment_rows.pop(key)
            local_assignments.append(a)
            for field, label in [("peptide_sequence", "Sequence"), ("fragment_type", "Frag Type")]:
                assert a[field] == row[cols[label]]
            for field, label in [("start_1based", "Start AA"), ("end_1based", "End AA"),
                                 ("observed_mass", "Observed Mass"), ("theoretical_mass", "Theoretical Mass"),
                                 ("mass_error_ppm", "Error")]:
                equal_number(a[field], row[cols[label]])
            assert a["parent_sequence"] == previous["parent_sequence"]
            assert a["parent_sequence"][a["start_1based"] - 1:a["end_1based"]] == a["peptide_sequence"]
            assert a["span_status"] == "matches_source_substrate"
            assert json.loads(a["modifications_json"]) == {h: raw_text(row[cols[h]]) for h in ("Localized Mod", "Unlocalized Mod", "Term Mod")}
            q = quantities[a["measurement_id"]]
            equal_number(q["quantitative_value"], row[cols["Intensity"]])
            assert q["peptide_sequence"] == a["peptide_sequence"]
            expected_quantities.add(q["measurement_id"])
            equal_number(a["n_assignments_reported"], row[cols.get("MultiMass Assignment", cols.get("MultiMass Ass"))])
            assert a["source_extra_identifier"] == (raw_text(row[7]) if shift else None)
            counts["n_assignment_rows"] += 1
        multiplicities = collections.Counter(a["native_mass_id_normalized"] for a in local_assignments)
        for a in local_assignments:
            assert a["n_assignments_replayed"] == multiplicities[a["native_mass_id_normalized"]]
            m = local_inputs.get(a["native_mass_id_normalized"])
            assert a["input_mass_id"] == (None if m is None else m["input_mass_id"])
            expected_status = "input_ids_not_supplied_no_inferred_link" if no_ids else (
                "input_id_not_found" if m is None else (
                    "input_id_and_intensity_agree" if quantities[a["measurement_id"]]["quantitative_value"] == quantities[m["measurement_id"]]["quantitative_value"] else "input_intensity_conflict"))
            assert a["mass_link_status"] == expected_status
    workbook.close()
assert not mass_rows and not assignment_rows
assert counts["n_input_mass_rows"] == 17862 and counts["n_assignment_rows"] == 8138
assert expected_quantities == {k for k, q in quantities.items() if q["quantity_kind"] in {"input_mass_sum_intensity", "assignment_intensity"}}

# Original inventory identities and public joined context, with no specimen inference.
files = json.loads((RAW / "files-project-route.json").read_text()) + json.loads((RAW / "files-page1.json").read_text())
raw_names = {r["fileName"] for r in files if r.get("fileCategory", {}).get("value") == "RAW"}
assert raw_names == {a["raw_file"] for a in acquisitions.values()} and len(raw_names) == 102
assert not any("Spleen20S_30min_R3" in f for f in raw_names)
for a in acquisitions.values():
    assert a["specimen_id"] is None and a["aliquot_id"] is None
    match = re.search(r"_([0-9]+)min_R([123])\.RAW$", a["raw_file"])
    if match:
        assert a["time_minutes"] == float(match[1]) and a["replicate_label_raw"] == match[2]
        has_substrate = "_AlphaSyn_" in a["raw_file"]
        has_enzyme = "20S_" in a["raw_file"]
        assert a["substrate_species"] == ("Homo sapiens" if has_substrate else None)
        assert a["enzyme_species"] == ("Mus musculus" if has_enzyme else None)
        organ = next((x for x in ("Brain", "Liver", "Spleen") if x + "20S_" in a["raw_file"]), None)
        assert a["organ_label"] == (organ or ("Liver" if has_enzyme else None))
    else:
        assert a["time_minutes"] is None
joined = query("assignments", select=["assignment_id", "quantity.quantitative_value", "acquisition.raw_file", "acquisition.time_minutes", "input_mass.native_mass_id_raw"], joins=["quantity", "acquisition", "input_mass"])
assert len(joined) == 8138
for row in joined:
    a = assignment_ids[row["assignment_id"]]
    assert row["quantity.quantitative_value"] == quantities[a["measurement_id"]]["quantitative_value"]
    view = views[a["view_id"]]
    assert row["acquisition.raw_file"] is not None, view
    label = view["sheet_name"]
    match = re.fullmatch(r"20S_AlphaSyn_(\d+)min_R([123])", label)
    if match:
        suffix = f"_AlphaSyn_20S_{match[1]}min_R{match[2]}.RAW"
    elif re.fullmatch(r"(?:AlphaSyn|20S)_\d+min_R[123]", label):
        suffix = "_" + label + ".RAW"
    else:
        match = re.fullmatch(r"(Brain|Liver|Spleen)_20S_aSyn_(\d+)min_([123])", label)
        assert match, label
        suffix = f"_AlphaSyn_{match[1]}20S_{match[2]}min_R{match[3]}.RAW"
    assert row["acquisition.raw_file"].endswith(suffix)
counts["n_joined_assignments"] = len(joined)

# Use the earlier independently derived, source-hashed publisher identity map.
supplements = {x["observation_id"]: x for x in query("supplement_rows")}
comparisons = query("reconciliations")
publisher = {records[supplements[x["source_subject_id"]]["source_record_id"]]: x for x in comparisons
             if x["assertion_predicate"] == "publisher_assignment_reconciliation"}
for table in previous["published_tables"]:
    source_file = "F2_SF6_SearchResults.xlsx" if table["sheet"] == "Kinetics" else "F3_SF7_SearchResults.xlsx"
    for expected in table["row_links"]:
        row = publisher.pop((table["file"], table["sheet"], expected["row"]))
        assignment = assignment_ids[row["target_subject_id"]]
        assert records[assignment["source_record_id"]] == (source_file, expected["source_sheet"], expected["source_rows"][0])
        detail = json.loads(row["details_json"])
        assert detail["intensity_equal"] == expected["quantity_equal"]
        assert detail["intensity_used_to_choose_identity"] is False
        assert len(detail["structural_candidates"]) == len(expected["structural_identity_source_rows"])
        counts["n_publisher_links"] += 1
        counts["n_publisher_quantity_conflicts"] += not detail["intensity_equal"]
assert not publisher and counts["n_publisher_links"] == 3427
assert counts["n_publisher_quantity_conflicts"] == 1
reused = [r for r in comparisons if r["assertion_predicate"] == "repeated_acquisition_source_views"]
assert len(reused) == len(previous["F1_F2_shared_sheets"]) == 9
assert all(r["assertion_status"] == "same_values" for r in reused)
assert {views[r["source_subject_id"]]["sheet_name"] for r in reused} == {r["sheet"] for r in previous["F1_F2_shared_sheets"]}
contexts = query("material_contexts")
assert len(contexts) == 103 and all(c["specimen_id"] is None and c["aliquot_id"] is None for c in contexts)
assert len(query("material_links")) == 102
result = {"passed": True, "counts": dict(counts), "workbooks": workbooks,
          "n_reused_source_views": len(reused), "n_acquisitions": len(acquisitions),
          "n_material_contexts": len(contexts), "public_query_join_checks": True,
          "producer_helpers_used_for_expected_values": False,
          "original_xml_reader_sha256": digest(UPSTREAM / "code/inventory.py"),
          "prior_independent_publisher_map_sha256": digest(UPSTREAM / "results/quantity-audit-layout-qualified.json"),
          "source_manifest_sha256": digest(Path(CONFIG["package"]) / "manifest.json"),
          "wall_seconds": time.monotonic() - START,
          "peak_rss_bytes": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss,
          "training_or_biological_prediction": None}
(ROOT / "results/independent-replay.json").write_text(json.dumps(result, indent=2) + "\n")
print(json.dumps(result), flush=True)
