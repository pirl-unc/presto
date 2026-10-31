"""Read-only source metadata and value-state census, without label admission."""

import csv
import gzip
import hashlib
import json
import math
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CACHE = Path("/private/tmp/presto-digestion-assets")


def dump(name, value):
    (ROOT / "results" / name).write_text(json.dumps(value, indent=2) + "\n")


def main():
    file = CACHE / "Identifier.xlsx"
    wb = openpyxl.load_workbook(file, read_only=True, data_only=False)
    sheets = {
        s.title: {
            "max_row": s.max_row,
            "max_column": s.max_column,
            "rows": [list(r) for r in s.iter_rows(values_only=True)],
        }
        for s in wb
    }
    dump(
        "identifier-raw-cells.json",
        {"sha256": hashlib.sha256(file.read_bytes()).hexdigest(), "sheets": sheets},
    )
    rows = sheets["Main_searches"]["rows"][1:]
    runs = {r[3].removesuffix(".raw"): r for r in rows if r[3]}
    assert len(runs) == 388 == len(rows)
    profiles = []
    for protein in sorted({r[0] for r in rows}):
        selected = [r for r in rows if r[0] == protein]
        profiles.append(
            {
                "source_name": protein,
                "runs": len(selected),
                "digest_times": sorted({r[1] for r in selected}),
                "biological_replicates": sorted({r[2] for r in selected}),
                "instruments": sorted({r[9] for r in selected}),
            }
        )
    dump("run-metadata-profiles.json", profiles)
    states = []
    for path in sorted((CACHE / "original-layers").glob("*/*finalKinetics.csv")):
        counts = Counter()
        repeated = quantified = 0
        with path.open(newline="") as source:
            for row in csv.DictReader(source):
                if row["biological_replicate"] == "NA":
                    counts["annotation_only_rows"] += 1
                    continue
                quantified += 1
                times = row["digestTimes"].split(";")
                repeated += len(times) != len(set(times))
                for value in row["intensities"].split(";"):
                    if value in ("NA", "NaN", ""):
                        counts["literal_" + repr(value)] += 1
                    elif not math.isfinite(float(value)):
                        counts["other_nonfinite"] += 1
                    else:
                        counts[
                            "zero"
                            if float(value) == 0
                            else "positive"
                            if float(value) > 0
                            else "negative"
                        ] += 1
        states.append(
            {
                "archive": path.parent.name,
                "quantified_rows": quantified,
                "rows_with_repeated_time": repeated,
                "values": dict(counts),
            }
        )
    dump("final-value-states.json", states)
    conflicts = []
    for number, row in enumerate(sheets["Contaminant_DB_generation"]["rows"][1:], 2):
        name = (row[1] or "").lower()
        enzyme = (row[2] or "").lower()
        if ("_try_" in name and enzyme == "chymotrypsin") or (
            "_chymo_" in name and enzyme == "trypsin"
        ):
            conflicts.append({"sheet_row": number, "cells": row})
    dump("enzyme-label-conflicts.json", conflicts)
    fastas = {}
    for path in (CACHE / "substrate-fastas").glob("*.fasta"):
        lines = path.read_text().splitlines()
        assert sum(x.startswith(">") for x in lines) == 1
        fastas[path.name] = "".join(x.strip() for x in lines if not x.startswith(">"))
    matches = []
    for path in sorted((CACHE / "original-layers").glob("*/*finalKinetics.csv")):
        with path.open(newline="") as source:
            sequences = {r["substrateSeq"] for r in csv.DictReader(source)}
        matches.append(
            {
                "archive": path.parent.name,
                "sequence_count": len(sequences),
                "matching_fasta_names": [
                    n for n, seq in fastas.items() if seq in sequences
                ],
                "sequence_sha256": [
                    hashlib.sha256(seq.encode()).hexdigest()
                    for seq in sorted(sequences)
                ],
            }
        )
    dump("construct-matches.json", matches)
    raw_census = []
    for path in sorted((CACHE / "original-layers").glob("*/QUANTITIES_raw.csv.gz")):
        n = 0
        sources = set()
        peptides = set()
        aliases = Counter()
        unknown = set()
        conflicts = []
        with gzip.open(path, "rt", newline="") as source:
            for row in csv.DictReader(source):
                n += 1
                sources.add(row["source"])
                peptides.add(row["pepSeq"])
                metadata = runs.get(row["source"])
                if metadata is None:
                    unknown.add(row["source"])
                    continue
                aliases[str(metadata[2]), row["biological_replicate"]] += 1
                if float(row["digestTime"]) != metadata[1]:
                    conflicts.append(
                        {
                            "source": row["source"],
                            "quantity_time": row["digestTime"],
                            "metadata_time": metadata[1],
                        }
                    )
        raw_census.append(
            {
                "archive": path.parent.name,
                "quantity_rows": n,
                "sources": sorted(sources),
                "unique_peptides": len(peptides),
                "unknown_sources": sorted(unknown),
                "time_conflicts": conflicts,
                "replicate_aliases": [
                    {"metadata": a, "quantity": b, "rows": v}
                    for (a, b), v in sorted(aliases.items())
                ],
            }
        )
    dump("raw-run-crosswalk.json", raw_census)
    doc = CACHE / "PSP_proteins_PRIDE_v2_JL.docx"
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(doc) as archive:
        tree = ET.fromstring(archive.read("word/document.xml"))
    txt = "\n".join(
        "".join(n.text or "" for n in p.findall(".//w:t", ns))
        for p in tree.findall(".//w:p", ns)
    )
    (ROOT / "results/deposited-methods.txt").write_text(txt)
    totals = Counter()
    for row in states:
        totals.update(row["values"])
    summary = {
        "main_search_runs": len(runs),
        "metadata_named_conditions": len(profiles),
        "raw_quantity_rows": sum(x["quantity_rows"] for x in raw_census),
        "unknown_raw_sources": sorted(
            {s for x in raw_census for s in x["unknown_sources"]}
        ),
        "time_conflict_rows": sum(len(x["time_conflicts"]) for x in raw_census),
        "quantified_final_rows": sum(x["quantified_rows"] for x in states),
        "final_value_states": dict(totals),
        "repeated_time_rows": sum(x["rows_with_repeated_time"] for x in states),
        "scope": "Literal source census; no training eligibility or biological missingness inference",
    }
    dump("metadata-audit.json", summary)
    print(json.dumps(summary), flush=True)


if __name__ == "__main__":
    main()
