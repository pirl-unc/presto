"""Preserve exact metadata locators for two manually reviewed source leads."""

import argparse
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    assert not args.out.exists()
    cases = []
    for name in ("project-download.xlsx", "sample-download.xlsx"):
        book = openpyxl.load_workbook(args.data_dir / name, read_only=True)
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        for row_index, row in enumerate(rows, start=2):
            if row[0] not in {"PPD017", "PPD070"}:
                continue
            cases.append(
                {
                    "workbook": name,
                    "sheet": sheet.title,
                    "row": row_index,
                    "fields": {
                        header: {"address": f"{get_column_letter(i)}{row_index}", "value": value}
                        for i, (header, value) in enumerate(zip(headers, row, strict=True), start=1)
                    },
                }
            )
        book.close()
    args.out.write_text(json.dumps(cases, indent=2, ensure_ascii=False) + "\n")
    print(f"Preserved {len(cases)} source metadata rows; no training labels.")


if __name__ == "__main__":
    main()
