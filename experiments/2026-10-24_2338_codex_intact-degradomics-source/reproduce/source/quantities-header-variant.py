"""Audit original quantities/lineage without selecting a supervised cohort."""

import collections
import csv
import json
import math
import re
import resource
import time

from inventory import OUT, RAW, check_resources, sha
from openpyxl import load_workbook


def finite_positive(value):
    return isinstance(value, (int, float)) and math.isfinite(value) and value > 0


def assignment_count(record):
    # One original sheet abbreviates this header; preserve both literal source keys.
    return record.get("MultiMass Assignment", record.get("MultiMass Ass"))


def rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator)
            yield sheet.title, header, iterator
    finally:
        workbook.close()


def main():
    start = time.monotonic()
    destination = OUT / "quantity-audit.json"
    assert not destination.exists()
    inventory = json.loads((OUT / "workbook-inventory.json").read_text())
    assert inventory["all_cells_match"]
    sources = {}
    summaries = []
    source_sequences = set()
    for path in sorted(RAW.glob("F*_SearchResults.xlsx")):
        for name, header, iterator in rows(path):
            check_resources()
            inputs, assignments = {}, []
            indexed = {v: i for i, v in enumerate(header) if v is not None}
            for row_number, row in enumerate(iterator, 2):
                if row[0] is not None:
                    assert row[0] not in inputs, (path.name, name, "duplicate input Mass ID")
                    inputs[row[0]] = {
                        "row": row_number,
                        "mass": row[1],
                        "adjusted_mass": row[2],
                        "intensity": row[3],
                    }
                if "Sequence" in indexed and row[indexed["Sequence"]] is not None:
                    assignment = {k: row[v] for k, v in indexed.items() if v >= indexed["Mass ID"]}
                    assignment["row"] = row_number
                    assignments.append(assignment)
                    if assignment["Start AA"] == 1 and assignment["End AA"] == 140:
                        source_sequences.add(assignment["Sequence"])
            counts = collections.Counter(x["Mass ID"] for x in assignments)
            sources[path.name, name] = {"inputs": inputs, "assignments": assignments}
            missing_ids = [x["row"] for x in assignments if x["Mass ID"] not in inputs]
            q_disagreements = [
                x["row"]
                for x in assignments
                if x["Mass ID"] in inputs and x["Intensity"] != inputs[x["Mass ID"]]["intensity"]
            ]
            count_disagreements = [
                x["row"] for x in assignments if assignment_count(x) != counts[x["Mass ID"]]
            ]
            summaries.append(
                {
                    "file": path.name,
                    "sheet": name,
                    "input_mass_rows": len(inputs),
                    "assignment_rows": len(assignments),
                    "unassigned_mass_ids": len(set(inputs) - set(counts)),
                    "assigned_mass_ids": len(counts),
                    "unique_assignment_mass_ids": sum(n == 1 for n in counts.values()),
                    "multi_assignment_mass_ids": sum(n > 1 for n in counts.values()),
                    "missing_input_ids": missing_ids,
                    "assignment_input_intensity_disagreements": q_disagreements,
                    "reported_assignment_count_disagreements": count_disagreements,
                    "nonpositive_or_nonnumeric_input_intensities": sum(
                        not finite_positive(x["intensity"]) for x in inputs.values()
                    ),
                    "input_intensity_sum": sum(
                        x["intensity"] for x in inputs.values() if finite_positive(x["intensity"])
                    ),
                }
            )
    assert len(source_sequences) == 1
    parent = next(iter(source_sequences))
    positional_mismatches = []
    for (filename, sheet), record in sources.items():
        for a in record["assignments"]:
            if parent[int(a["Start AA"]) - 1 : int(a["End AA"])] != a["Sequence"]:
                positional_mismatches.append([filename, sheet, a["row"]])

    # Source identities determine joins. Quantity values never select a match.
    published = []
    mappings = {}
    for suffix, sheet, source in [
        ("sd04", "Kinetics", "F2_SF6_SearchResults.xlsx"),
        ("sd05", "Organs", "F3_SF7_SearchResults.xlsx"),
    ]:
        path = RAW / f"pnas.2419607122.{suffix}.xlsx"
        records, counts = [], collections.Counter()
        for name, header, iterator in rows(path):
            if name != sheet:
                continue
            for number, row in enumerate(iterator, 2):
                if row[0] is None:
                    continue
                record = dict(zip(header, row, strict=True))
                sample = record["Sample ID"]
                if suffix == "sd04":
                    time_label, rep = re.fullmatch(r"([0-9]+min)_Rep([123])", sample).groups()
                    source_sheet = f"20S_AlphaSyn_{time_label}_R{rep}"
                else:
                    organ, time_label, rep = re.fullmatch(
                        r"(Brain|Liver|Spleen)_([0-9]+min)_([123])", sample
                    ).groups()
                    source_sheet = f"{organ}_20S_aSyn_{time_label}_{rep}"
                mappings[f"{suffix}:{sample}"] = [source, source_sheet]
                assignments = sources.get((source, source_sheet), {}).get("assignments", [])
                frag = {
                    "N-term": "B Fragment",
                    "C-term": "Y Fragment",
                    "Internal": "BY Int Fragment",
                }[record["Frag Type"]]
                matches = [
                    x
                    for x in assignments
                    if all(x[k] == record[k] for k in ["Sequence", "Start AA", "End AA"])
                    and x["Frag Type"] == frag
                ]
                status = (
                    "missing_source_sheet"
                    if (source, source_sheet) not in sources
                    else (
                        "unique_source_identity"
                        if len(matches) == 1
                        else "ambiguous_or_missing_identity"
                    )
                )
                quantity_equal = (
                    len(matches) == 1 and matches[0]["Intensity"] == record["Intensity"]
                )
                compare_keys = [
                    "Observed Mass",
                    "Theoretical Mass",
                    "Error",
                    "Intensity",
                    "Formula",
                ]
                mismatched_fields = [
                    k for k in compare_keys if len(matches) == 1 and matches[0][k] != record[k]
                ]
                counts[status] += 1
                counts["quantity_equal"] += quantity_equal
                counts["positive_finite_intensity"] += finite_positive(record["Intensity"])
                counts["single_reported_assignment"] += len(matches) == 1 and (
                    assignment_count(matches[0]) == 1
                )
                records.append(
                    {
                        "row": number,
                        "sample": sample,
                        "source_sheet": source_sheet,
                        "start": record["Start AA"],
                        "end": record["End AA"],
                        "sequence": record["Sequence"],
                        "status": status,
                        "source_rows": [x["row"] for x in matches],
                        "quantity_equal": quantity_equal,
                        "mismatched_fields": mismatched_fields,
                    }
                )
        by_sequence = collections.defaultdict(set)
        by_sample = collections.Counter()
        for record in records:
            by_sequence[record["sequence"]].add(record["sample"])
            by_sample[record["sample"]] += 1
        published.append(
            {
                "file": path.name,
                "sheet": sheet,
                "rows": len(records),
                "counts": dict(counts),
                "sample_rows": dict(sorted(by_sample.items())),
                "unique_sequences": len(by_sequence),
                "length_min": min(map(len, by_sequence)),
                "length_max": max(map(len, by_sequence)),
                "sequences_at_multiple_samples": sum(len(v) > 1 for v in by_sequence.values()),
                "row_links": records,
            }
        )

    reuse = []
    for (filename, sheet), record in sources.items():
        if filename.startswith("F1") and ("F2_SF6_SearchResults.xlsx", sheet) in sources:
            other = sources["F2_SF6_SearchResults.xlsx", sheet]

            def normalized(records):
                return collections.Counter(
                    json.dumps({k: v for k, v in row.items() if k != "row"}, sort_keys=True)
                    for row in records
                )

            reuse.append(
                {
                    "sheet": sheet,
                    "input_masses_equal": record["inputs"] == other["inputs"],
                    "assignment_multiset_equal": normalized(record["assignments"])
                    == normalized(other["assignments"]),
                    "input_rows": len(record["inputs"]),
                    "assignment_rows": len(record["assignments"]),
                }
            )
    check_resources()
    result = {
        "parent_sequence": parent,
        "parent_length": len(parent),
        "source_sheets": summaries,
        "positional_mismatches": positional_mismatches,
        "published_tables": published,
        "sample_mappings": mappings,
        "F1_F2_shared_sheets": reuse,
        "seconds": time.monotonic() - start,
        "peak_rss_bytes": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss,
        "inventory_sha256": sha(OUT / "workbook-inventory.json"),
        "training_admission": False,
    }
    destination.write_text(json.dumps(result, indent=2) + "\n")
    with (OUT / "source-sheet-summary.csv").open("x") as stream:
        writer = csv.DictWriter(stream, fieldnames=summaries[0], lineterminator="\n")
        writer.writeheader()
        writer.writerows(summaries)
    print(
        json.dumps(
            {
                "source_sheets": len(summaries),
                "parent_length": len(parent),
                "positional_mismatches": positional_mismatches,
                "F1_F2_shared_sheets": reuse,
                "published_tables": [
                    {k: v for k, v in x.items() if k != "row_links"} for x in published
                ],
                "seconds": result["seconds"],
                "peak_rss_bytes": result["peak_rss_bytes"],
            }
        )
    )


if __name__ == "__main__":
    main()
